"""data/hackathons.xlsx. Events is the only data sheet; Dashboard, My Events
and Review Queue are rebuilt from it on every write, and edits there are not
read back. Formatting uses conditional rules, so colours follow the user's
sorting, filtering and dropdown choices."""

from __future__ import annotations

import re
from collections import Counter
from datetime import date
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from tracker import config, normalize
from tracker import events as ev

SHEETS = ("Dashboard", "Events", "My Events", "Review Queue", "Lists")
STATUS_COLOURS = {  # value: (fill, text)
    "Registration open": ("C6EFCE", "006100"), "Upcoming": ("DDEBF7", "1F4E79"),
    "Registration not yet open": ("FFF2CC", "7F6000"),
    "Applications required": ("FCE4D6", "843C0C"), "Waitlist": ("FCE4D6", "843C0C"),
    "Date TBC": ("EDEDED", "404040"), "Recurring / annual - date TBC": ("EDEDED", "404040"),
    "Registration closed": ("FFC7CE", "9C0006"), "Completed": ("F2F2F2", "808080"),
}
MY_STATUS_COLOURS = {
    "Interested": ("E4DFEC", "3F3151"), "Registered": ("C6EFCE", "006100"),
    "Applied": ("DDEBF7", "1F4E79"), "Accepted": ("A9D08E", "1E4620"),
    "Attended": ("EDEDED", "404040"), "Not going": ("F2F2F2", "808080"),
}
URGENCY = ("FFC7CE", "FCE4D6", "FFF2CC")  # paired with config deadline_bands
WRAP = {"name", "action", "description", "eligibility", "notes", "my_notes", "evidence",
        "technical_area", "location", "organisation", "host", "prize", "sponsors",
        "discovery_method", "reasons"}
WIDTHS = {"id": 10, "my_status": 13, "name": 30, "status": 16, "action": 22,
          "registration_deadline": 12, "start_date": 12, "end_date": 12,
          "last_verified": 12, "days_to_deadline": 10, "days_to_event": 10,
          "technical_area": 18, "location": 18, "eligibility": 36, "evidence": 36,
          "description": 50, "notes": 36, "my_notes": 28, "official_url": 28,
          "source_url": 28, "organisation": 20, "host": 20, "discovery_method": 20,
          "reasons": 40}
MY_COLS = [("ID", "id"), ("My Status", "my_status"), ("Event Name", "name"),
           ("Action", "action"), ("Status", "status"),
           ("Registration Deadline", "registration_deadline"),
           ("Days Until Deadline", "days_to_deadline"), ("Start Date", "start_date"),
           ("Days Until Event", "days_to_event"), ("Location", "location"),
           ("Official Event URL", "official_url"), ("My Notes", "my_notes")]
REVIEW_COLS = [("ID", "id"), ("Event Name", "name"), ("Needs Checking", "reasons"),
               ("Status", "status"), ("Start Date", "start_date"),
               ("Official Event URL", "official_url"), ("Source URL", "source_url"),
               ("Last Verified", "last_verified")]
_CAP = 25  # rows per Dashboard block; the rest are on the Events sheet


def _fill(hex_: str) -> PatternFill:
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def _short(url: str) -> str:
    parts = urlsplit(url)
    text = parts.netloc.removeprefix("www.") + parts.path.rstrip("/")
    return text if len(text) <= 40 else text[:37] + "..."


def _read_user_cells(path) -> dict[str, dict]:
    if not path.exists():
        return {}
    wb = load_workbook(path, read_only=True, data_only=True)
    if "Events" not in wb.sheetnames:
        return {}
    it = wb["Events"].iter_rows(values_only=True)
    index = {name: i for i, name in enumerate(next(it, ())) if name}
    if "ID" not in index:
        return {}
    return {str(r[index["ID"]]): {k: r[index[label]] for label, k in ev.USER_OWNED.items()
                                  if label in index and r[index[label]] is not None}
            for r in it if r[index["ID"]]}


def _colour_by_value(ws, letter: str, last: int, colours: dict) -> None:
    for value, (fill, text) in colours.items():
        ws.conditional_formatting.add(f"{letter}2:{letter}{last}", CellIsRule(
            operator="equal", formula=[f'"{value}"'], fill=_fill(fill), font=Font(color=text)))


def _table(ws, cols, rows, cfg, freeze="A2") -> None:
    """One database-style table at A1: dark header, calm body, fixed row height
    so hundreds of rows stay scannable. The caller adds banding last, so every
    other rule outranks it."""
    font = cfg["font"]
    keys = [k for _, k in cols]
    ws.append([label for label, _ in cols])
    for r in rows:
        ws.append([normalize.excel_safe(v) if isinstance(v, str) else v
                   for v in (r.get(k) for k in keys)])
    body = Font(name=font, size=10)
    for i, key in enumerate(keys, 1):
        letter = get_column_letter(i)
        head = ws.cell(1, i)
        head.fill = _fill("7F6000" if key in ev.USER_OWNED.values() else "1F3864")
        head.font = Font(name=font, size=11, bold=True, color="FFFFFF")
        head.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[letter].width = WIDTHS.get(key, 13)
        wrap = key in WRAP
        for cell in ws[letter][1:]:
            cell.font = body
            cell.alignment = Alignment(wrap_text=wrap, vertical="top" if wrap else "center")
            if key in ev.DATE_FIELDS:
                cell.number_format = "yyyy-mm-dd"
            elif key.endswith("_url") and cell.value:
                cell.hyperlink, cell.value = cell.value, _short(cell.value)
                cell.font = Font(name=font, size=10, color="0563C1", underline="single")
    ws.row_dimensions[1].height = 32
    for n in range(2, ws.max_row + 1):
        ws.row_dimensions[n].height = 45
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions


def _band(ws) -> None:
    if ws.max_row > 1:
        ws.conditional_formatting.add(
            f"A2:{get_column_letter(ws.max_column)}{ws.max_row}",
            FormulaRule(formula=["MOD(ROW(),2)=0"], fill=_fill("F5F7FA")))


def _style_status_columns(ws, cols, cfg) -> None:
    keys = [k for _, k in cols]
    last = max(ws.max_row, 2)
    letter = lambda k: get_column_letter(keys.index(k) + 1)  # noqa: E731
    if "status" in keys:
        _colour_by_value(ws, letter("status"), last, STATUS_COLOURS)
    if "my_status" in keys:
        _colour_by_value(ws, letter("my_status"), last, MY_STATUS_COLOURS)
    if "days_to_deadline" in keys:
        col, low = letter("days_to_deadline"), 0
        for top, fill in zip(cfg["deadline_bands"], URGENCY):
            ws.conditional_formatting.add(f"{col}2:{col}{last}", FormulaRule(
                formula=[f"AND(ISNUMBER({col}2),{col}2>={low},{col}2<={top})"],
                fill=_fill(fill)))
            low = top + 1


def _protect(ws) -> None:
    """Views are regenerated; protection (no password) says so, while sort,
    filter and column resizing keep working."""
    ws.protection.sheet = True
    ws.protection.sort = ws.protection.autoFilter = ws.protection.formatColumns = False


def write(store: list[dict], today: date, leads: int = 0) -> bool:
    """Rebuild every sheet from the store. The user's cells are read back by ID
    first and win; an open workbook is skipped whole, never half-written."""
    path = config.EVENTS_WORKBOOK
    if (path.parent / f"~${path.name}").exists():
        return False
    cfg = config.load("events")
    mine = _read_user_cells(path)
    rows = sorted((ev.display(e, today, mine.get(e["id"])) for e in store),
                  key=lambda r: ev.sort_key(r, today))
    wb = load_workbook(path) if path.exists() else Workbook()
    for name in (*SHEETS, "Sheet"):
        if name in wb.sheetnames:
            del wb[name]
    sheets = {name: wb.create_sheet(name, i) for i, name in enumerate(SHEETS)}
    _dashboard(sheets["Dashboard"], rows, today, cfg, leads)
    _events(sheets["Events"], rows, cfg)
    _my_events(sheets["My Events"], rows, today, cfg)
    _review_queue(sheets["Review Queue"], rows, cfg)
    _lists(sheets["Lists"], cfg)
    wb.active = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
    except PermissionError:
        return False
    return True


def _events(ws, rows, cfg) -> None:
    _table(ws, ev.COLUMNS, rows, cfg, freeze="D2")
    _style_status_columns(ws, ev.COLUMNS, cfg)
    options = cfg["my_statuses"]
    dv = DataValidation(type="list", formula1=f"=Lists!$A$1:$A${len(options)}",
                        allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B2:B{max(ws.max_row, 2)}")
    _band(ws)


def _lists(ws, cfg) -> None:
    for n, value in enumerate(cfg["my_statuses"], 1):
        ws.cell(n, 1, value)
    ws.sheet_state = "hidden"


def _my_events(ws, rows, today, cfg) -> None:
    mine = sorted((r for r in rows if r.get("my_status")),
                  key=lambda r: (ev.next_date(r, today), r["name"].casefold()))
    _table(ws, MY_COLS, mine, cfg)
    _style_status_columns(ws, MY_COLS, cfg)
    _band(ws)
    _protect(ws)


def _review_queue(ws, rows, cfg) -> None:
    queue = [{**r, "reasons": "; ".join(ev.review_reasons(r))} for r in rows
             if r["status"] != "Completed" and ev.review_reasons(r)]
    _table(ws, REVIEW_COLS, queue, cfg)
    _style_status_columns(ws, REVIEW_COLS, cfg)
    _band(ws)
    _protect(ws)


def _dashboard(ws, rows, today, cfg, leads) -> None:
    """Plain two-column blocks, computed here from the same rows as Events,
    so the numbers can never drift from the table."""
    font, week = cfg["font"], cfg["deadline_bands"][1]
    upcoming = [r for r in rows if r["status"] != "Completed"]
    due = lambda r, n: r["days_to_deadline"] is not None and r["days_to_deadline"] <= n  # noqa: E731
    registered = [r for r in upcoming if r.get("my_status") in ("Registered", "Accepted")]
    interested = [r for r in upcoming if r.get("my_status") == "Interested"]
    action = [r for r in upcoming if r["action"] and (
        (due(r, week) and r.get("my_status") not in ("Registered", "Accepted", "Applied", "Not going"))
        or (r.get("my_status") == "Interested" and r["status"] == "Registration open")
        or (r in registered and r["days_to_event"] is not None and r["days_to_event"] <= week))]
    split = lambda text: [p.strip() for p in re.split(r"[,/;]", text or "") if p.strip()]  # noqa: E731
    areas = Counter(a for r in upcoming for a in split(r["technical_area"]))
    methods = Counter(m for r in rows for m in split(r["discovery_method"]))
    by_date = lambda rs: sorted(rs, key=lambda r: ev.next_date(r, today))  # noqa: E731
    n = 1

    def line(values, bold=False, size=10, link=None):
        nonlocal n
        for i, v in enumerate(values, 1):
            c = ws.cell(n, i, normalize.excel_safe(v) if isinstance(v, str) else v)
            c.font = Font(name=font, size=size, bold=bold)
            c.alignment = Alignment(vertical="center")
            if isinstance(v, date):
                c.number_format = "yyyy-mm-dd"
        if link:
            ws.cell(n, 1).hyperlink = link
            ws.cell(n, 1).font = Font(name=font, size=10, color="0563C1", underline="single")
        n += 1

    def block(title, head, items, cells):
        nonlocal n
        n += 1
        line([title], bold=True, size=12)
        line(head, bold=True)
        for c in range(1, len(head) + 1):
            ws.cell(n - 1, c).fill = _fill("D9E1F2")
        for r in items[:_CAP]:
            line(cells(r), link=r.get("official_url") if isinstance(r, dict) else None)
        if len(items) > _CAP:
            line([f"...and {len(items) - _CAP} more - see the Events sheet"])
        if not items:
            line(["None"])

    line(["UK Hackathon Tracker"], bold=True, size=14)
    line([f"Updated {today.isoformat()} - generated from the Events sheet; edit there, not here"])
    block("At a glance", ["Measure", "Count"], [
        ("Total upcoming events", len(upcoming)),
        ("Registration open", sum(r["status"] == "Registration open" for r in upcoming)),
        ("Deadlines in next 7 days", sum(due(r, 7) for r in upcoming)),
        ("Deadlines in next 14 days", sum(due(r, 14) for r in upcoming)),
        ("My registered events", len(registered)),
        ("My interested events", len(interested)),
        ("Need review", sum(bool(ev.review_reasons(r)) for r in upcoming)),
        ("Open leads (found, not yet verified)", leads),
    ], list)
    block("Action needed", ["Event", "Action", "Deadline", "Days", "My Status"], by_date(action),
          lambda r: [r["name"], r["action"], r["registration_deadline"],
                     r["days_to_deadline"], r.get("my_status")])
    block("Upcoming deadlines (next 14 days)", ["Event", "Deadline", "Days", "Status", "My Status"],
          by_date([r for r in upcoming if due(r, 14)]),
          lambda r: [r["name"], r["registration_deadline"], r["days_to_deadline"],
                     r["status"], r.get("my_status")])
    block("My registered events", ["Event", "Start", "Days", "Location"], by_date(registered),
          lambda r: [r["name"], r["start_date"], r["days_to_event"], r["location"]])
    block("My interested events", ["Event", "Status", "Deadline", "Action"], by_date(interested),
          lambda r: [r["name"], r["status"], r["registration_deadline"], r["action"]])
    block("Events by UK region", ["Region", "Upcoming"],
          Counter(r["region"] for r in upcoming).most_common(), list)
    block("Events by technical area", ["Area", "Upcoming"], areas.most_common(), list)
    block("Events by discovery method", ["Method", "All events"], methods.most_common(), list)
    for letter, width in zip("ABCDE", (38, 24, 12, 16, 14)):
        ws.column_dimensions[letter].width = width
    _protect(ws)

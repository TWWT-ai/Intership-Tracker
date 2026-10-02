"""data/hackathons.xlsx: user cells survive, the layout reads well, views are derived."""

import json
from datetime import date, datetime

from click.testing import CliRunner
from openpyxl import load_workbook

from main import cli
from tests.conftest import make_event as cand
from tracker import config, events, events_excel

T = date(2026, 9, 30)


def build(*cands, today=T):
    events.ingest(list(cands), today)
    assert events_excel.write(events.load_store(), today)
    return load_workbook(config.EVENTS_WORKBOOK)


def rows(wb, sheet="Events"):
    ws = wb[sheet]
    head = [c.value for c in ws[1]]
    return [dict(zip(head, (c.value for c in r))) for r in ws.iter_rows(min_row=2)]


def set_my_status(wb, values: dict[int, str]):
    ws = wb["Events"]
    head = [c.value for c in ws[1]]
    for n, value in values.items():
        ws.cell(n, head.index("My Status") + 1, value)
    wb.save(config.EVENTS_WORKBOOK)


# --- Events sheet ----------------------------------------------------------

def test_my_status_survives_a_rebuild():
    wb = build(cand())
    ws = wb["Events"]
    head = [c.value for c in ws[1]]
    ws.cell(2, head.index("My Status") + 1, "Registered")
    ws.cell(2, head.index("My Notes") + 1, "go with Sam")
    wb.save(config.EVENTS_WORKBOOK)
    r = rows(build(cand(prize="GBP 1000", last_verified="2026-10-07"),
                   today=date(2026, 10, 7)))[0]
    assert (r["My Status"], r["My Notes"], r["Prize"]) == ("Registered", "go with Sam", "GBP 1000")


def test_upcoming_sorts_before_completed():
    wb = build(cand(name="Old Hack", organisation="X", start_date="2026-01-10",
                    official_url="https://old.example/"), cand())
    assert [r["Event Name"] for r in rows(wb)] == ["HackSheffield 11", "Old Hack"]


def test_dates_are_dates_and_formulas_are_defused():
    r = rows(build(cand(description='=HYPERLINK("http://x")')))[0]
    assert isinstance(r["Start Date"], datetime) and r["End Date"] is None
    assert not str(r["Event Description"]).startswith("=")


def test_evidence_and_discovery_method_are_shown():
    r = rows(build(cand()))[0]
    assert r["Format Evidence"].startswith("Teams of up to four")
    assert r["Discovery Method"] == "Society Search"


def test_url_shows_short_text_but_links_in_full():
    long = "https://hacksheffield.co.uk/events/2026/hacksheffield-eleven/register?ref=x"
    ws = build(cand(official_url=long))["Events"]
    head = [c.value for c in ws[1]]
    cell = ws.cell(2, head.index("Official Event URL") + 1)
    assert cell.hyperlink.target == long and len(cell.value) <= 40


def test_status_colour_is_a_rule_and_text_stays():
    wb = build(cand(status="Registration open"))
    assert rows(wb)[0]["Status"] == "Registration open"
    formulas = [r.formula for cf in wb["Events"].conditional_formatting for r in cf.rules]
    assert ['"Registration open"'] in formulas


def test_no_merged_cells_anywhere():
    assert all(not ws.merged_cells.ranges for ws in build(cand()).worksheets)


def test_busy_workbook_writes_nothing():
    config.EVENTS_WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    (config.EVENTS_WORKBOOK.parent / "~$hackathons.xlsx").write_text("")
    assert events_excel.write([], T) is False
    assert not config.EVENTS_WORKBOOK.exists()


# --- views -----------------------------------------------------------------

def dash(wb, label):
    return next(r[1].value for r in wb["Dashboard"].iter_rows() if r[0].value == label)


def test_dashboard_counts_come_from_the_data():
    wb = build(cand(status="Registration open", registration_deadline="2026-10-04"),
               cand(name="Bath Hack", organisation="Bath Hack", official_url="https://bathhack.com/",
                    start_date="2027-02-20", registration_deadline="2026-10-12",
                    discovery_method="University Search"),
               cand(name="Old Hack", organisation="X", official_url="https://old.example/",
                    start_date="2026-01-10"))
    assert wb.active.title == "Dashboard"
    assert dash(wb, "Total upcoming events") == 2
    assert dash(wb, "Registration open") == 1
    assert dash(wb, "Deadlines in next 7 days") == 1
    assert dash(wb, "Deadlines in next 14 days") == 2
    assert dash(wb, "Society Search") == 2 and dash(wb, "University Search") == 1


def test_my_events_lists_only_rows_with_my_status_in_date_order():
    wb = build(cand(), cand(name="Bath Hack", organisation="Bath Hack",
                            official_url="https://bathhack.com/", start_date="2026-10-20"),
               cand(name="Quiet Hack", organisation="Q", official_url="https://q.example/"))
    set_my_status(wb, {2: "Registered", 3: "Registered"})
    assert events_excel.write(events.load_store(), T)
    names = [r["Event Name"] for r in rows(load_workbook(config.EVENTS_WORKBOOK), "My Events")]
    assert names == ["Bath Hack", "HackSheffield 11"]


def test_review_queue_keeps_uncertain_events():
    wb = build(cand(student_eligible=""), cand(name="Bath Hack", organisation="Bath Hack",
                                               official_url="https://bathhack.com/"))
    q = rows(wb, "Review Queue")
    assert [r["Event Name"] for r in q] == ["HackSheffield 11"]
    assert q[0]["Needs Checking"] == "Student eligibility unclear"
    assert len(rows(wb)) == 2  # still in Events too


# --- CLI -------------------------------------------------------------------

def test_cli_ingest_writes_workbook(tmp_path):
    f = tmp_path / "c.json"
    f.write_text(json.dumps([cand()]))
    res = CliRunner().invoke(cli, ["events", "--ingest", str(f)])
    assert res.exit_code == 0, res.output
    assert "added 1" in res.output and config.EVENTS_WORKBOOK.exists()


def test_cli_dry_run_saves_nothing(tmp_path):
    f = tmp_path / "c.json"
    f.write_text(json.dumps([cand()]))
    res = CliRunner().invoke(cli, ["events", "--ingest", str(f), "--dry-run"])
    assert "new 1" in res.output and not config.EVENTS_STORE.exists()


def test_cli_plan_writes_the_sweep_plan():
    res = CliRunner().invoke(cli, ["events", "--plan"])
    assert res.exit_code == 0, res.output
    assert len(json.loads(config.EVENTS_PLAN.read_text())) == 11


def test_cli_render_with_no_events_writes_nothing():
    res = CliRunner().invoke(cli, ["events"])
    assert res.exit_code == 0 and not config.EVENTS_WORKBOOK.exists()

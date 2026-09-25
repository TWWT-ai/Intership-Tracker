"""Company careers boards, against canned platform responses. No network."""

import json

import pytest
from openpyxl import load_workbook

from sources import careers
from sources.fetch import SourceError
from tests.conftest import make_job
from tracker import config, dedupe, excel, normalize

SETTINGS = config.load("sources")["sources"]["careers"]


class FakeResponse:
    def __init__(self, payload, status=200):
        self.text = json.dumps(payload) if not isinstance(payload, str) else payload
        self.status = status

    @property
    def ok(self):
        return 200 <= self.status < 300


class RoutedSession:
    """Answers each URL from a table; records what was asked for."""

    def __init__(self, routes: dict):
        self.routes = routes
        self.calls: list[tuple[str, str]] = []
        self.http = self
        self.limiter = self

    def configure(self, host, seconds):
        pass

    def _answer(self, method, url):
        self.calls.append((method, url))
        for prefix, payload in self.routes.items():
            if url.startswith(prefix):
                if isinstance(payload, Exception):
                    raise payload
                return payload if isinstance(payload, FakeResponse) else FakeResponse(payload)
        return FakeResponse({}, status=404)

    def get(self, url, **kwargs):
        return self._answer("GET", url)

    def post_json(self, url, payload, **kwargs):
        return self._answer("POST", url + "#offset=" + str(payload["offset"]))


def settings_with(**boards):
    return {**SETTINGS, "discover_from_sources": False,
            "boards": [{"company": c, "url": u} for c, u in boards.items()]}


GREENHOUSE = {"jobs": [
    {"id": 4001111, "title": "Software Engineering Intern - Summer 2027",
     "location": {"name": "London"}, "absolute_url": "https://boards.greenhouse.io/acme/jobs/4001111",
     "departments": [{"name": "Engineering"}], "content": "&lt;p&gt;Python and C++&lt;/p&gt;"},
    {"id": 4002222, "title": "Software Engineering Intern", "location": {"name": "New York, NY"},
     "absolute_url": "https://boards.greenhouse.io/acme/jobs/4002222"},
    {"id": 4003333, "title": "Senior Software Engineer", "location": {"name": "London"},
     "absolute_url": "https://boards.greenhouse.io/acme/jobs/4003333"},
    {"id": 4004444, "title": "Finance Intern", "location": {"name": "London"},
     "absolute_url": "https://boards.greenhouse.io/acme/jobs/4004444", "departments": [{"name": "Finance"}]},
    {"id": 4005555, "title": "Summer Intern", "location": {"name": "Cambridge, UK"},
     "absolute_url": "https://boards.greenhouse.io/acme/jobs/4005555", "departments": [{"name": "Engineering"}]},
]}


class TestGreenhouse:
    @pytest.fixture
    def jobs(self):
        session = RoutedSession({"https://boards-api.greenhouse.io/v1/boards/acme/jobs": GREENHOUSE})
        return careers.search_jobs(session, settings_with(Acme="https://boards.greenhouse.io/acme"))

    def test_keeps_only_uk_student_tech_roles(self, jobs):
        """US office, senior role and a non-tech internship are all dropped."""
        assert sorted(j.sources[0].external_id for j in jobs) == [
            "greenhouse:acme:4001111", "greenhouse:acme:4005555"]

    def test_a_tech_department_counts_when_the_title_is_generic(self, jobs):
        assert any(j.job_title == "Summer Intern" for j in jobs)

    def test_records_where_it_was_found(self, jobs):
        job = next(j for j in jobs if j.sources[0].external_id.endswith(":4001111"))
        assert job.primary_source == "careers"
        assert job.sources[0].via == "Greenhouse"
        assert job.company == "Acme"
        assert job.application_url == "https://boards.greenhouse.io/acme/jobs/4001111"

    def test_description_is_unescaped_text(self, jobs):
        job = next(j for j in jobs if j.sources[0].external_id.endswith(":4001111"))
        assert "Python and C++" in job.description and "<p>" not in job.description

    def test_work_type_only_when_the_posting_says(self, jobs):
        by_id = {j.sources[0].external_id: j for j in jobs}
        assert by_id["greenhouse:acme:4001111"].work_type == "Summer Internship"


class TestOtherPlatforms:
    def test_lever_eu_board_uses_the_eu_api(self):
        session = RoutedSession({"https://api.eu.lever.co/v0/postings/acme": [
            {"id": "0a1b2c3d-1111-2222-3333-444455556666", "text": "Data Engineering Placement",
             "categories": {"location": "London", "team": "Data", "commitment": "Intern"},
             "hostedUrl": "https://jobs.eu.lever.co/acme/0a1b2c3d-1111-2222-3333-444455556666",
             "applyUrl": "https://jobs.eu.lever.co/acme/0a1b2c3d-1111-2222-3333-444455556666/apply",
             "descriptionPlain": "SQL"}]})
        jobs = careers.search_jobs(session, settings_with(Acme="https://jobs.eu.lever.co/acme/x"))
        assert len(jobs) == 1 and jobs[0].work_type == "Industrial Placement"
        assert jobs[0].application_url.endswith("/apply")

    def test_lever_requirement_lists_are_part_of_the_posting(self):
        """Lever keeps "What We Require" apart from the description; dropping it
        lost every "Graduating in 2028" a Lever board ever stated."""
        session = RoutedSession({"https://api.lever.co/v0/postings/acme": [
            {"id": "0a1b2c3d-1111-2222-3333-444455556666", "text": "Software Engineer, Internship",
             "categories": {"location": "London", "commitment": "Internship"},
             "hostedUrl": "https://jobs.lever.co/acme/0a1b2c3d-1111-2222-3333-444455556666",
             "descriptionPlain": "Build things with customers.",
             "lists": [{"text": "What We Require",
                        "content": "<li>Proficiency in Python.</li><li>Graduating in 2028.</li>"}],
             "additionalPlain": "Offer deadline: two weeks."}]})
        postings = careers._lever(session, careers.board_from_url("https://jobs.lever.co/acme", "Acme"), {}, None)
        text = postings[0]["description"]
        assert "What We Require" in text and "Graduating in 2028." in text and "Build things" in text

    def test_ashby_employment_type_marks_a_student_role(self):
        session = RoutedSession({"https://api.ashbyhq.com/posting-api/job-board/acme": {"jobs": [
            {"id": "11111111-2222-3333-4444-555555555555", "title": "Software Engineer",
             "location": "London", "employmentType": "Intern", "isListed": True,
             "jobUrl": "https://jobs.ashbyhq.com/acme/11111111-2222-3333-4444-555555555555"},
            {"id": "99999999-2222-3333-4444-555555555555", "title": "Software Engineer",
             "location": "London", "employmentType": "FullTime", "isListed": True,
             "jobUrl": "https://jobs.ashbyhq.com/acme/99999999-2222-3333-4444-555555555555"}]}})
        jobs = careers.search_jobs(session, settings_with(Acme="https://jobs.ashbyhq.com/acme"))
        assert [j.sources[0].via for j in jobs] == ["Ashby"]

    def test_workable(self):
        session = RoutedSession({"https://apply.workable.com/api/v1/widget/accounts/acme": {"jobs": [
            {"shortcode": "ABC123DEF", "title": "Machine Learning Intern", "city": "Oxford",
             "country": "United Kingdom", "url": "https://apply.workable.com/acme/j/ABC123DEF/",
             "description": "<p>PyTorch</p>"}]}})
        jobs = careers.search_jobs(session, settings_with(Acme="https://apply.workable.com/acme/"))
        assert len(jobs) == 1 and jobs[0].location

    def test_workday_searches_then_opens_only_screened_postings(self):
        api = "https://acme.wd3.myworkdayjobs.com/wday/cxs/acme/Careers"
        session = RoutedSession({
            f"{api}/jobs": {"total": 2, "jobPostings": [
                {"title": "Technology Summer Intern 2027", "externalPath": "/job/London/Tech-Intern_R1"},
                {"title": "Chemistry Intern", "externalPath": "/job/Bintulu/Chem-Intern_R2"}]},
            f"{api}/job/London/Tech-Intern_R1": {"jobPostingInfo": {
                "title": "Technology Summer Intern 2027", "location": "London",
                "country": {"descriptor": "United Kingdom"}, "jobReqId": "R1",
                "externalUrl": "https://acme.wd3.myworkdayjobs.com/Careers/job/London/Tech-Intern_R1",
                "jobDescription": "<p>Java</p>"}},
        })
        settings = {**settings_with(Acme="https://acme.wd3.myworkdayjobs.com/en-US/Careers/job/x"),
                    "workday_queries": ["intern"]}
        jobs = careers.search_jobs(session, settings)

        assert [j.sources[0].external_id for j in jobs] == ["workday:acme:wd3:Careers:R1"]
        opened = [u for m, u in session.calls if m == "GET"]
        assert opened == [f"{api}/job/London/Tech-Intern_R1"]     # the chemistry role never opened


class TestResilience:
    def test_a_failing_board_reuses_its_last_good_postings(self):
        good = RoutedSession({"https://boards-api.greenhouse.io/v1/boards/acme/jobs": GREENHOUSE,
                              "https://api.ashbyhq.com/posting-api/job-board/beta": {"jobs": []}})
        settings = settings_with(Acme="https://boards.greenhouse.io/acme",
                                 Beta="https://jobs.ashbyhq.com/beta")
        first = careers.search_jobs(good, settings)
        config.ensure_dirs()
        (config.SNAPSHOT_DIR / "careers.json").write_text(
            json.dumps({"jobs": [j.to_dict() for j in first]}), encoding="utf-8")

        flaky = RoutedSession({"https://boards-api.greenhouse.io/v1/boards/acme/jobs": SourceError("down"),
                               "https://api.ashbyhq.com/posting-api/job-board/beta": {"jobs": []}})
        again = careers.search_jobs(flaky, settings)
        assert len(again) == len(first) == 2

    def test_every_board_failing_is_a_source_failure(self):
        session = RoutedSession({"https://boards-api.greenhouse.io": SourceError("down")})
        with pytest.raises(SourceError):
            careers.search_jobs(session, settings_with(Acme="https://boards.greenhouse.io/acme"))


class TestDiscovery:
    def test_boards_are_found_in_other_sources_links(self):
        config.ensure_dirs()
        trackr_job = make_job(company="Acme", application_url="https://boards.greenhouse.io/acme/jobs/4001111")
        (config.SNAPSHOT_DIR / "trackr.json").write_text(
            json.dumps({"jobs": [trackr_job.to_dict()]}), encoding="utf-8")

        boards = careers.discover_boards({**SETTINGS, "boards": []})
        assert [(b.key, b.company) for b in boards] == [("greenhouse:acme", "Acme")]

    def test_an_unrecognised_configured_board_is_skipped(self):
        assert careers.discover_boards(settings_with(Acme="https://acme.example.com/careers")) == []


class TestPostingKey:
    @pytest.mark.parametrize("a,b", [
        ("https://boards.greenhouse.io/acme/jobs/8726366002?gh_jid=8726366002",
         "https://job-boards.eu.greenhouse.io/acme/jobs/8726366002"),
        ("https://www.acme.com/careers/job?gh_jid=8726366002",
         "https://boards.greenhouse.io/acme/jobs/8726366002"),
        ("https://acme.wd3.myworkdayjobs.com/en-US/Careers/job/London/SWE-Intern_R123?source=Trackr",
         "https://acme.wd3.myworkdayjobs.com/Careers/job/London/SWE-Intern_R123"),
        ("https://jobs.lever.co/acme/0a1b2c3d-1111-2222-3333-444455556666/apply",
         "https://jobs.lever.co/acme/0a1b2c3d-1111-2222-3333-444455556666"),
    ])
    def test_one_posting_in_different_shapes_has_one_key(self, a, b):
        assert normalize.posting_key(a) and normalize.posting_key(a) == normalize.posting_key(b)

    def test_a_generic_careers_page_has_no_key(self):
        """No posting id means no claim of sameness."""
        assert normalize.posting_key("https://www.acme.com/careers/") == ""


class TestMergingWithTrackr:
    def _pair(self):
        trackr_job = make_job(company="Acme", title="2027 Summer Internship Programme", location="United Kingdom",
                              source="trackr", external_id="trk1",
                              application_url="https://boards.greenhouse.io/acme/jobs/4001111?utm_source=Trackr",
                              work_type="Summer Internship")
        session = RoutedSession({"https://boards-api.greenhouse.io/v1/boards/acme/jobs": GREENHOUSE})
        site_jobs = careers.search_jobs(session, settings_with(Acme="https://boards.greenhouse.io/acme"))
        return trackr_job, site_jobs

    def test_same_posting_becomes_one_row_despite_title_and_location(self):
        trackr_job, site_jobs = self._pair()
        merged, count = dedupe.deduplicate([trackr_job] + site_jobs)
        assert count == 1
        row = next(j for j in merged if "trackr" in j.source_names())
        assert row.source_names() == ["careers", "trackr"]

    def test_sources_cell_names_the_platform(self):
        trackr_job, site_jobs = self._pair()
        merged, _ = dedupe.deduplicate([trackr_job] + site_jobs)
        dedupe.assign_ids(merged)
        ws = load_workbook(excel.write(merged))[excel.SHEET_JOBS]
        col = excel.col_index()
        values = {ws.cell(r, col["Sources"]).value for r in range(2, ws.max_row + 1)}
        assert "Trackr, Company site (Greenhouse)" in values
        assert "Company site (Greenhouse)" in values


def search_settings(**boards):
    return {**settings_with(**boards), "search_queries": ["intern"], "search_max_pages": 1,
            "platforms": list(careers.PLATFORM_NAMES)}


class TestSiblingProgrammes:
    """Programmes an aggregator never listed, found on the employer's own board."""

    @pytest.fixture
    def jobs(self):
        session = RoutedSession({"https://boards-api.greenhouse.io/v1/boards/acme/jobs": {"jobs": [
            {"id": 7001001, "title": "Technology Development Programme", "location": {"name": "London"},
             "absolute_url": "https://boards.greenhouse.io/acme/jobs/7001001",
             "departments": [{"name": "Early Careers"}],
             "content": "&lt;p&gt;Eligibility: second-year students on a four-year degree "
                        "with a placement year.&lt;/p&gt;"},
            {"id": 7002002, "title": "Graduate Development Programme - Technology",
             "location": {"name": "London"}, "absolute_url": "https://boards.greenhouse.io/acme/jobs/7002002",
             "content": "First-year students welcome."},
            {"id": 7003003, "title": "Spring Insight Programme 2027", "location": {"name": "London"},
             "absolute_url": "https://boards.greenhouse.io/acme/jobs/7003003",
             "departments": [{"name": "Early Careers"}]},
            {"id": 7004004, "title": "Technology Programme", "location": {"name": "London"},
             "absolute_url": "https://boards.greenhouse.io/acme/jobs/7004004", "content": "Python."},
        ]}})
        found = careers.search_jobs(session, settings_with(Acme="https://boards.greenhouse.io/acme"))
        return {j.sources[0].external_id.rsplit(":", 1)[1]: j for j in found}

    def test_a_programme_named_otherwise_is_kept_by_who_may_apply(self, jobs):
        job = jobs["7001001"]
        assert job.work_type == "Industrial Placement"
        assert "placement year" in job.year_of_study

    def test_a_firm_wide_spring_programme_is_kept_without_a_tech_word(self, jobs):
        assert jobs["7003003"].work_type == "Spring Week"

    def test_graduate_programmes_and_programmes_with_no_evidence_are_dropped(self, jobs):
        assert "7002002" not in jobs and "7004004" not in jobs


ORACLE = "https://acme.fa.em2.oraclecloud.com"
ORACLE_API = f"{ORACLE}/hcmRestApi/resources/latest"


class TestOracle:
    @pytest.fixture
    def run(self):
        session = RoutedSession({
            f"{ORACLE_API}/recruitingCEJobRequisitions?": {"items": [{"TotalJobsCount": 3, "requisitionList": [
                {"Id": "5001", "Title": "Technology Discovery Programme 2027",
                 "PrimaryLocation": "London, England, United Kingdom", "PrimaryLocationCountry": "GB"},
                {"Id": "5002", "Title": "Global Markets Summer Analyst", "PrimaryLocation": "London"},
                {"Id": "5003", "Title": "Graduate Development Programme - Technology",
                 "PrimaryLocation": "London"}]}]},
            f"{ORACLE_API}/recruitingCEJobRequisitionDetails?expand=all&onlyData=true&finder=ById;Id=%225001%22": {
                "items": [{"Id": "5001", "Title": "Technology Discovery Programme 2027",
                           "ExternalDescriptionStr": "<p>Who can apply</p><ul><li>You are a first-year "
                                                     "undergraduate student on any degree.</li></ul>"}]},
        })
        board = f"{ORACLE}/hcmUI/CandidateExperience/en/sites/CX_1/job/4000/?utm_source=Trackr"
        return session, careers.search_jobs(session, search_settings(Acme=board))

    def test_finds_the_spring_week_by_its_requirements(self, run):
        _, jobs = run
        assert len(jobs) == 1
        job = jobs[0]
        assert job.sources[0].external_id == "oracle:acme.fa.em2.oraclecloud.com:CX_1:5001"
        assert job.sources[0].via == "Oracle"
        assert job.work_type == "Spring Week" and "first-year" in job.year_of_study
        assert job.application_url == f"{ORACLE}/hcmUI/CandidateExperience/en/sites/CX_1/job/5001"

    def test_opens_only_screened_requisitions(self, run):
        session, _ = run
        details = [u for _, u in session.calls if "RequisitionDetails" in u]
        assert len(details) == 1 and "5001" in details[0]

    def test_merges_with_the_aggregators_link_to_the_same_requisition(self, run):
        _, jobs = run
        trackr_link = f"{ORACLE}/hcmUI/CandidateExperience/en/sites/CX_1/job/5001/?utm_medium=tracker"
        assert normalize.posting_key(trackr_link) == normalize.posting_key(jobs[0].application_url) != ""


class TestEightfold:
    MS = "https://apply.careers.microsoft.com"

    def test_microsoft_board_is_found_from_its_old_share_link(self):
        board = careers.board_from_url("https://jobs.careers.microsoft.com/global/en/share/1871538/?x=1", "Microsoft")
        assert board.key == "eightfold:apply.careers.microsoft.com:microsoft.com"

    def test_the_shared_host_without_a_domain_names_no_board(self):
        assert careers.board_from_url("https://app.eightfold.ai/careers?query=x", "Qualcomm") is None

    def test_searches_then_opens_only_screened_positions(self):
        session = RoutedSession({
            f"{self.MS}/api/pcsx/search": {"data": {"positions": [
                {"id": 1970393556991724, "name": "Cambridge Internship Program - Robotics Systems and Control",
                 "locations": ["United Kingdom, Cambridgeshire, Cambridge"], "department": "Research Sciences"},
                {"id": 1970393556991730, "name": "Postdoctoral Researcher (Foundational AI)",
                 "locations": ["United Kingdom, Cambridgeshire, Cambridge"]}]}},
            f"{self.MS}/api/pcsx/position_details?position_id=1970393556991724": {"data": {
                "id": 1970393556991724, "name": "Cambridge Internship Program - Robotics Systems and Control",
                "locations": ["United Kingdom, Cambridgeshire, Cambridge"],
                "jobDescription": "<ul><li>Open to students in the penultimate year of their degree.</li>"
                                  "<li>Contract Type: Internship (12 Weeks)</li></ul>"}},
        })
        jobs = careers.search_jobs(session, search_settings(
            Microsoft="https://jobs.careers.microsoft.com/global/en/share/1871538/"))

        assert [j.sources[0].external_id for j in jobs] == [
            "eightfold:apply.careers.microsoft.com:microsoft.com:1970393556991724"]
        assert jobs[0].work_type == "Summer Internship"
        assert jobs[0].application_url == f"{self.MS}/careers/job/1970393556991724"
        opened = [u for _, u in session.calls if "position_details" in u]
        assert len(opened) == 1

    def test_a_tenant_answering_with_a_web_page_is_a_failed_board(self):
        session = RoutedSession({"https://aexp.eightfold.ai/api/pcsx/search":
                                 FakeResponse("<!DOCTYPE html><html></html>")})
        with pytest.raises(SourceError):
            careers.search_jobs(session, search_settings(
                AmEx="https://aexp.eightfold.ai/careers/job/39765754"))


class TestAmazon:
    def test_keeps_uk_tech_internships_without_guessing_a_season(self):
        session = RoutedSession({"https://www.amazon.jobs/en/search.json": {"hits": 2, "jobs": [
            {"id_icims": "3072061", "title": "2026 Software Dev Engineer Intern - UK",
             "normalized_location": "London, England, GBR", "job_category": "Software Development",
             "is_intern": True, "job_path": "/en/jobs/3072061/2026-software-dev-engineer-intern-uk",
             "description": "Build services.", "basic_qualifications": "Enrolled in a degree"},
            {"id_icims": "10487050", "title": "Operations & Logistics Internship, Amazon Logistics",
             "normalized_location": "Belfast, Northern Ireland, GBR", "job_category": "Administrative Support",
             "is_intern": True, "job_path": "/en/jobs/10487050/operations-logistics-internship"}]}})
        trackr_link = "https://www.amazon.jobs/en/jobs/3072061/2026-software-dev-engineer-intern?utm_source=Trackr"
        jobs = careers.search_jobs(session, search_settings(Amazon=trackr_link))

        assert [j.sources[0].external_id for j in jobs] == ["amazon:www.amazon.jobs:3072061"]
        assert jobs[0].work_type == ""
        assert normalize.posting_key(trackr_link) == normalize.posting_key(jobs[0].application_url)


class TestOnePostingOneRow:
    def test_one_requisition_on_two_sites_of_a_tenant_is_one_row(self):
        """BNY lists each requisition on both its CX_1001 and BNY-Careers sites."""
        session = RoutedSession({
            f"{ORACLE_API}/recruitingCEJobRequisitions?": {"items": [{"TotalJobsCount": 1, "requisitionList": [
                {"Id": "6001", "Title": "Software Engineering Summer Internship 2027",
                 "PrimaryLocation": "London, United Kingdom"}]}]},
            f"{ORACLE_API}/recruitingCEJobRequisitionDetails?": {"items": [
                {"Id": "6001", "Title": "Software Engineering Summer Internship 2027"}]},
        })
        site = f"{ORACLE}/hcmUI/CandidateExperience/en/sites"
        settings = {**search_settings(Acme=f"{site}/Careers/job/6001"),
                    "boards": [{"company": "Acme", "url": f"{site}/Careers/job/6001"},
                               {"company": "Acme", "url": f"{site}/CX_1/job/6001"}]}
        jobs = careers.search_jobs(session, settings)
        assert [j.sources[0].external_id for j in jobs] == ["oracle:acme.fa.em2.oraclecloud.com:CX_1:6001"]


import re  # noqa: E402
from pathlib import Path  # noqa: E402

FIXTURES = Path(__file__).parent / "fixtures"


def fixture(name):
    return (FIXTURES / name).read_text(encoding="utf-8")


class TestOleeo:
    BOARD = ("https://blackrock.tal.net/vx/lang-en-GB/mobile-0/brand-3/user-1735377/xf-d85d6b29451d/"
             "candidate/jobboard/vacancy/1/adv/")
    TRACKR_LINK = ("https://blackrock.tal.net/vx/lang-en-GB/mobile-0/brand-3/user-1735377/xf-bce35dd61046/"
                   "candidate/so/pm/1/pl/1/opp/12223-2027-Spring-Insight-Event-EMEA/en-GB?utm_source=Trackr")

    def _session(self):
        board_html = fixture("oleeo_blackrock_board.html")
        spring = re.search(r'href="([^"]*/opp/12223-[^"]*)"', board_html).group(1)
        return RoutedSession({spring: FakeResponse(fixture("oleeo_blackrock_spring.html")),
                              self.BOARD: FakeResponse(board_html)})

    def test_board_is_found_from_a_posting_link(self):
        board = careers.board_from_url(self.TRACKR_LINK, "BlackRock")
        assert board.key == "oleeo:blackrock.tal.net:1"
        assert board.url == ("https://blackrock.tal.net/vx/lang-en-GB/mobile-0/brand-3/user-1735377/"
                             "xf-bce35dd61046/candidate/jobboard/vacancy/1/adv/")

    def test_finds_the_spring_insight_event_and_its_graduation_year(self):
        jobs = careers.search_jobs(self._session(), search_settings(BlackRock=self.BOARD))
        assert [j.job_title for j in jobs] == ["2027 Spring Insight Event - EMEA"]
        job = jobs[0]
        assert job.work_type == "Spring Week"
        assert job.expected_graduation == "2029"
        assert any(city in job.location for city in ("London", "Edinburgh"))
        assert job.sources[0].via == "Oleeo"

    def test_rows_outside_the_uk_are_never_opened(self):
        session = self._session()
        careers.search_jobs(session, search_settings(BlackRock=self.BOARD))
        opened = [u for _, u in session.calls if "/opp/" in u]
        assert opened and not any(re.search(r"AMRS|APAC|Mexico", u) for u in opened)

    def test_same_posting_as_the_aggregators_link(self):
        jobs = careers.search_jobs(self._session(), search_settings(BlackRock=self.BOARD))
        assert normalize.posting_key(self.TRACKR_LINK) == normalize.posting_key(jobs[0].application_url) != ""

    def test_paging_follows_next_page(self):
        board = ("https://morganstanley.tal.net/vx/lang-en-GB/mobile-0/brand-2/user-5764832/xf-8c7b6a6cce83/"
                 "candidate/jobboard/vacancy/1/adv/")
        session = RoutedSession({f"{board}?start=50": FakeResponse(fixture("oleeo_morganstanley_board_p2.html")),
                                 board: FakeResponse(fixture("oleeo_morganstanley_board.html"))})
        careers.search_jobs(session, {**search_settings(MorganStanley=board), "oleeo_max_pages": 2})
        pages = [u for _, u in session.calls if "/jobboard/" in u]
        assert pages == [board, f"{board}?start=50"]

    def test_an_anti_bot_check_skips_the_board_without_going_further(self):
        board = ("https://bankcampuscareers.tal.net/vx/lang-en-GB/mobile-0/brand-4/user-4698662/xf-3fa8353e9d47/"
                 "candidate/jobboard/vacancy/1/adv/")
        session = RoutedSession({board: FakeResponse(fixture("oleeo_bankofamerica_challenge.html"))})
        with pytest.raises(SourceError):
            careers.search_jobs(session, search_settings(BankOfAmerica=board))
        assert [u for _, u in session.calls if "/opp/" in u] == []


class TestYourGraduationYear:
    def test_a_uk_posting_for_your_year_is_kept_whatever_its_team(self, isolated_data):
        config.PROFILE.write_text(json.dumps({"graduation_year": 2029}), encoding="utf-8")
        session = RoutedSession({"https://boards-api.greenhouse.io/v1/boards/acme/jobs": {"jobs": [
            {"id": 8001001, "title": "Markets Discovery Week", "location": {"name": "London"},
             "absolute_url": "https://boards.greenhouse.io/acme/jobs/8001001", "departments": [{"name": "Sales"}],
             "content": "&lt;p&gt;For students graduating in 2029.&lt;/p&gt;"},
            {"id": 8002002, "title": "Markets Discovery Week", "location": {"name": "London"},
             "absolute_url": "https://boards.greenhouse.io/acme/jobs/8002002", "departments": [{"name": "Sales"}],
             "content": "&lt;p&gt;For students graduating in 2028.&lt;/p&gt;"},
        ]}})
        jobs = careers.search_jobs(session, settings_with(Acme="https://boards.greenhouse.io/acme"))
        assert [(j.sources[0].external_id, j.work_type, j.expected_graduation) for j in jobs] == [
            ("greenhouse:acme:8001001", "Spring Week", "2029")]


class TestFirmWideProgrammes:
    """A programme with a brand name is the tech intake too, so it is read. A
    title naming another function or field is that intake's, so it is not."""
    wanted = careers._Filter(SETTINGS)

    @pytest.mark.parametrize("title,opened", [
        ("2027 EMEA Future Leaders Spring Insight Programme", True),
        ("Spring Insight Programme 2027", True),
        ("Ignite Undergraduate Programme", True),
        ("Technology Summer Intern 2027", True),
        ("Finance Industrial Placement Programme", False),
        ("Sales Summer Internship", False),
        ("Chemistry Intern", False),
        ("Mechanical Engineering Placement Year", False),
    ])
    def test_a_branded_programme_is_opened_but_another_function_is_not(self, title, opened):
        assert self.wanted.screen(title, "Blackstone") is opened


class TestLocation:
    wanted = careers._Filter(SETTINGS)

    @pytest.mark.parametrize("raw,shown", [
        ("Abu Dhabi, Amsterdam, Edinburgh, London; United Arab Emirates, Netherlands, United Kingdom", "Edinburgh"),
        ("London, England, United Kingdom", "London, England, United Kingdom"),
        ("Cambridge, MA; London", "London"),
        ("New York, NY", "New York, NY"),
        ("Sydney, New South Wales, Australia", "Sydney, New South Wales, Australia"),
    ])
    def test_the_uk_office_is_the_one_shown(self, raw, shown):
        assert careers._first_uk_location(raw, self.wanted) == shown

    @pytest.mark.parametrize("raw,uk", [
        ("Sydney, New South Wales, Australia", False),
        ("Boston, New England", False),
        ("Cardiff, Wales", True),
        ("Edinburgh, Scotland; Sydney, New South Wales", True),
    ])
    def test_a_place_abroad_named_after_a_uk_country_is_not_the_uk(self, raw, uk):
        assert self.wanted.is_uk(raw) is uk

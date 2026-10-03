"""The assessment applied to the sheet, safe page reading, coverage, and the
read-only dry run. Offline: pages come from a fake session that records every
request, so "one GET, nothing else" is checked, not assumed."""
import hashlib
import json
from datetime import date, datetime, timedelta

import pytest
from openpyxl import load_workbook

from scripts import calendar_sync
from sources.fetch import Response
from tests.conftest import make_job
from tracker import config, dedupe, excel, pipeline, relevance, role_filter, verify
from tracker import deadline_extract as dx
from tracker.models import ApplicationStatus, OpenStatus

TODAY = date.today()
SOON = TODAY + timedelta(days=10)
PAST = TODAY - timedelta(days=5)


def rows(*jobs):
    dedupe.assign_ids(list(jobs))
    return list(jobs)


def sheet_ids():
    ws = load_workbook(config.WORKBOOK)[excel.SHEET_JOBS]
    col = [c.value for c in ws[1]].index("ID") + 1
    return {ws.cell(r, col).value for r in range(2, ws.max_row + 1)}


@pytest.fixture
def on(monkeypatch):
    monkeypatch.setattr(role_filter, "enabled", lambda: True)


def technical(job):
    job.role_relevance, job.study_level, job.degree_subject_fit = (
        relevance.TECHNICAL, relevance.CURRENT_UNDERGRADUATE, relevance.CS_ACCEPTED)
    return job


# ---------------------------------------------------------------- off by default

def test_with_the_assessment_off_nothing_changes():
    assert role_filter.enabled() is False
    keep, other = rows(make_job(company="Keep", title="Data Intern"), make_job(company="Other", title="Audit Intern"))
    other.role_relevance = relevance.NOT_TECHNICAL
    excel.write([keep, other])
    assert sheet_ids() == {keep.id, other.id}
    assert excel.SHEET_ASSESSMENT not in load_workbook(config.WORKBOOK).sheetnames


# ---------------------------------------------------------------- which rows the sheet shows

def test_firm_evidence_hides_unclear_and_unassessed_stay(on):
    tech, unclear, audit, unassessed = rows(
        technical(make_job(company="A", title="Software Intern")), make_job(company="B", title="Data Intern"),
        make_job(company="C", title="Audit Intern"), make_job(company="D", title="Finance Intern"))
    unclear.role_relevance, audit.role_relevance = relevance.UNCLEAR, relevance.NOT_TECHNICAL
    audit.role_relevance_reason = "duties are audit, tax work, not computing"
    excel.write([tech, unclear, audit, unassessed])
    assert sheet_ids() == {tech.id, unclear.id, unassessed.id}      # the title never hides "Finance Intern"
    ws = load_workbook(config.WORKBOOK)[excel.SHEET_ASSESSMENT]
    header = [c.value for c in ws[1]]
    first = {h: ws.cell(2, i + 1).value for i, h in enumerate(header)}
    assert (first["ID"], first["On Internships sheet?"], first["Role relevance"]) == (audit.id, "No", "NOT_TECHNICAL")
    assert "Deadline candidates" in header and "Eligible for You" in header


def test_a_passed_listed_deadline_stays_visible_and_is_marked_expired(on):
    listed = rows(make_job(company="L", title="Software Intern", deadline=PAST))[0]
    listed.deadline_status = dx.LISTED
    assert role_filter.decide(listed, TODAY)[0] == role_filter.SHOW
    assert "deadline passed - listed date, not confirmed by the employer" in role_filter.decide(listed, TODAY)[1]
    assert role_filter.deadline_marker(listed, TODAY) == "EXPIRED (listed date, not employer-confirmed)"


def test_a_passed_employer_confirmed_deadline_is_kept_and_marked_not_closed(on):
    job = rows(make_job(company="E", title="Software Intern", deadline=PAST))[0]
    job.deadline_status, job.deadline_confirmed = dx.CONFIRMED, PAST
    job.hiring_status, job.hiring_evidence = role_filter.hiring_status(job)
    assert role_filter.decide(job, TODAY)[0] == role_filter.SHOW
    assert role_filter.deadline_marker(job, TODAY) == "EXPIRED (employer-confirmed date)"
    assert job.hiring_status == role_filter.UNKNOWN                  # expiry is not closure


def test_expired_rows_sort_to_the_bottom_whatever_their_hiring_status(on):
    expired_open, upcoming, expired_unknown = rows(
        make_job(company="A", title="Old Open", deadline=PAST), make_job(company="B", title="Soon", deadline=SOON),
        make_job(company="C", title="Old Unknown", deadline=PAST - timedelta(days=3)))
    expired_open.hiring_status = role_filter.EMPLOYER_OPEN
    assert [j.company for j in excel.sort_jobs([expired_open, upcoming, expired_unknown])][0] == "B"
    excel.write([expired_open, upcoming, expired_unknown])
    ws = load_workbook(config.WORKBOOK)[excel.SHEET_JOBS]
    header = [c.value for c in ws[1]]
    assert header.index("Deadline Status") == header.index("Deadline") + 1
    assert header.index("Hiring Status") == header.index("Open Status") + 1
    marks = [ws.cell(r, header.index("Deadline Status") + 1).value for r in range(2, ws.max_row + 1)]
    assert marks[0] in ("", None) and all(str(m).startswith("EXPIRED") for m in marks[1:])


def test_an_explicitly_closed_posting_stays_at_the_bottom_marked_closed(on):
    closed, live = rows(make_job(company="E", title="Software Intern", deadline=SOON),
                        make_job(company="F", title="Software Intern", deadline=SOON + timedelta(days=30)))
    closed.closed_evidence = 'employer page says "applications are closed" (jobs.example.com)'
    decision, reason = role_filter.decide(closed, TODAY)
    assert decision == role_filter.SHOW and reason.startswith("CLOSED")
    assert role_filter.open_status(closed) is OpenStatus.CLOSED
    assert [j.company for j in excel.sort_jobs([closed, live])] == ["F", "E"]   # sorted to the bottom


def test_your_status_keeps_any_row_through_a_full_rebuild(on):
    audit, other = rows(make_job(company="C", title="Audit Intern", deadline=PAST),
                        make_job(company="D", title="Data Intern"))
    audit.role_relevance = relevance.NOT_TECHNICAL
    audit.deadline_status, audit.deadline_confirmed = dx.CONFIRMED, PAST
    excel.write([audit, other], status_updates={audit.id: ApplicationStatus.APPLIED})
    again = rows(make_job(company="C", title="Audit Intern", deadline=PAST), make_job(company="D", title="Data Intern"))
    again[0].role_relevance = relevance.NOT_TECHNICAL
    excel.write(again)                                                # the 01:00 rebuild path
    assert again[0].id == audit.id and sheet_ids() == {audit.id, other.id}
    state = excel.read_user_state()
    assert ApplicationStatus.parse(state[audit.id][excel.STATUS_COLUMN]) is ApplicationStatus.APPLIED


# ---------------------------------------------------------------- reading pages safely

POSTING = ("<html lang='en-GB'><body><main><h1>Software Engineering Placement</h1>"
           "<script type='application/ld+json'>{\"@type\": \"JobPosting\", \"validThrough\": \"2026-11-08\"}</script>"
           "<p>What you'll do: build software and write code in Python. " + "More about the role. " * 40 +
           "Applications close 08/11/2026.</p>"
           "<form><input name='email'><button>Apply now</button></form></main></body></html>")
FORM = ("<html lang='en-GB'><body><main><h1>Apply</h1><form><label>First name</label><input name='first_name'>"
        "<label>Last name</label><input name='last_name'><label>Email</label><input name='email'>"
        "<label>Upload your CV</label><input type='file' name='cv'><button type='submit'>Submit</button>"
        "</form></main></body></html>")


class FakeHttp:
    def __init__(self, pages):
        self.pages, self.calls = pages, []

    def get(self, url, **_):
        self.calls.append(url)
        page = self.pages[url]
        status, text = page if isinstance(page, tuple) else (200, page)
        return Response(url=url, status=status, text=text)


class FakeSession:
    def __init__(self, pages):
        self.http = FakeHttp(pages)

    @property
    def browser(self):
        pytest.fail("a browser was opened")


def test_a_public_posting_at_an_apply_url_is_read():
    url = "https://lbg.wd3.myworkdayjobs.com/en-GB/careers/job/Software-Placement_1/apply"
    session = FakeSession({url: POSTING})
    page = verify.read_page(session, url, config.load("verification"), allow_browser=False)
    assert page.kind == "posting" and page.locale == "en-GB" and page.valid_through == "2026-11-08"
    order = dx.order_from_locale(page.locale)
    assert [c.date for c in dx.candidates(page.text, source="employer page", url=url, method="employer page",
                                          employer_page=True, order=order)] == ["2026-11-08"]
    assert session.http.calls == [url]                               # one GET, nothing followed


def test_an_actual_application_form_stops_the_read():
    url = "https://jobs.example.com/careers/job/123"
    session = FakeSession({url: FORM})
    page = verify.read_page(session, url, config.load("verification"), allow_browser=False)
    assert page.kind == "application_step" and page.text == ""
    assert "stopped at application step" in page.note
    assert session.http.calls == [url]


WORKDAY_URL = "https://lbg.wd3.myworkdayjobs.com/en-US/careers/job/London/Software-Placement_1/apply"
WORKDAY_API = "https://lbg.wd3.myworkdayjobs.com/wday/cxs/lbg/careers/job/London/Software-Placement_1"


def workday_feed(description):
    return json.dumps({"jobPostingInfo": {"title": "Software Placement", "location": "London",
                                          "jobDescription": "<p>" + description + "</p>"}})


def test_a_workday_posting_is_read_through_its_job_feed_and_confirms_the_deadline(on):
    job = rows(make_job(company="Lloyds", title="Software Placement", deadline=date(2026, 11, 1),
                        application_url=WORKDAY_URL))[0]
    job.deadline_status = dx.LISTED
    text = "What you'll do: build software in Python. " + "More about the role. " * 30 + "Applications close 8 November 2026."
    session = FakeSession({WORKDAY_API: workday_feed(text)})
    outcome = verify.verify_assessed(session, job, config.load("verification"), set(), relevance.Listings([]), TODAY)
    assert outcome == "checked"
    assert job.deadline_confirmed == date(2026, 11, 8) and job.deadline == date(2026, 11, 8)
    assert job.deadline_status == dx.CONFLICT and job.deadline_provisional == date(2026, 11, 1)
    assert session.http.calls == [WORKDAY_API]                       # one request, the feed only


def test_an_en_us_link_does_not_make_a_uk_date_month_first(on):
    job = rows(make_job(company="Lloyds", title="Software Placement", application_url=WORKDAY_URL))[0]
    text = "What you'll do: build software in Python. " + "More about the role. " * 30 + "Applications close 03/10/2026."
    verify.verify_assessed(FakeSession({WORKDAY_API: workday_feed(text)}), job, config.load("verification"), set(),
                           relevance.Listings([]), TODAY)
    page = [c for c in job.deadline_evidence if c["method"] == "employer page"]
    assert [(c["date"], c["status"]) for c in page] == [("", dx.AMBIGUOUS)]


def test_a_posting_the_feed_no_longer_has_is_unavailable_not_closed(on):
    job = rows(make_job(company="Lloyds", title="Software Placement", deadline=PAST, application_url=WORKDAY_URL))[0]
    outcome = verify.verify_assessed(FakeSession({WORKDAY_API: (404, "")}), job, config.load("verification"), set(),
                                     relevance.Listings([]), TODAY)
    assert outcome == "checked" and job.hiring_status == role_filter.UNAVAILABLE
    assert job.closed_evidence == "" and job.open_status is not OpenStatus.CLOSED
    assert role_filter.decide(job, TODAY)[0] == role_filter.SHOW
    assert job.checked_url == WORKDAY_URL and job.last_checked


def test_a_script_built_page_is_not_counted_as_checked(on):
    url = "https://careers.example.com/job/123"
    job = rows(make_job(company="X", title="Software Intern", application_url=url))[0]
    shell = "<html lang='en-US'><body><div id='root'></div><script>load()</script></body></html>"
    outcome = verify.verify_assessed(FakeSession({url: shell}), job, config.load("verification"), set(),
                                     relevance.Listings([]), TODAY)
    assert outcome == "no_text" and job.deadline_confirmed is None


def test_a_passed_listed_date_never_closes_a_row_by_itself(on):
    url = "https://jobs.example.com/careers/job/9"
    job = rows(make_job(company="X", title="Software Intern", deadline=PAST, application_url=url))[0]
    page = POSTING.replace("Applications close 08/11/2026.", "").replace('"validThrough": "2026-11-08"', '"x": 1')
    verify.verify_assessed(FakeSession({url: page}), job, config.load("verification"), set(),
                           relevance.Listings([]), TODAY)
    assert job.open_status is not OpenStatus.CLOSED and job.sheet_decision == role_filter.SHOW


# ---------------------------------------------------------------- coverage and priority

def test_the_queue_puts_approaching_deadlines_then_stale_unknowns_first_and_rotates():
    settings = dict(config.load("verification"), budget={"max_verifications": 4, "revalidate_after_days": 3,
                                                          "always_verify": {"deadline_within_days": 14}})
    soon = make_job(company="Soon", title="A", deadline=SOON, application_url="https://jobs.a.com/job/1")
    stale = make_job(company="Stale", title="B", application_url="https://jobs.b.com/job/2")
    stale.last_checked = (datetime.now() - timedelta(days=10)).isoformat()
    fresh = [make_job(company=f"New{i}", title=f"C{i}", application_url=f"https://jobs.c{i}.com/job/{i}")
             for i in range(4)]
    recent = make_job(company="Recent", title="D", application_url="https://jobs.d.com/job/9")
    recent.last_checked, recent.hiring_status = datetime.now().isoformat(), role_filter.EMPLOYER_OPEN
    aggregator = make_job(company="Agg", title="E", application_url="https://www.gradcracker.com/hub/1/job")
    jobs = rows(*fresh, stale, soon, recent, aggregator)
    chosen, later = verify.select_assessed(jobs, settings, TODAY)
    assert [j.company for j in chosen[:2]] == ["Soon", "Stale"]
    assert len(chosen) == 4 and len(later) == 2                       # 6 due: the rest wait their turn
    assert all(j.check_outcome == "over_budget" for j in later)
    assert recent.check_outcome == "not_due" and aggregator.check_outcome == "no_employer_url"
    for j in chosen:                                                  # this run's checks...
        j.last_checked, j.hiring_status = datetime.now().isoformat(), role_filter.EMPLOYER_OPEN
    again, _ = verify.select_assessed(jobs, settings, TODAY)
    assert {j.company for j in again} == {j.company for j in later}   # ...so the next run resumes with the rest


def test_dry_run_changes_nothing_and_opens_nothing(monkeypatch, tmp_path):
    from sources import fetch

    monkeypatch.setattr(fetch, "Session", lambda *a, **k: pytest.fail("the dry run opened a session"))
    tech = make_job(company="Palantir", title="Software Engineer Intern", deadline=SOON,
                    description="What you'll do: " + "build software, write code in Python and Java, APIs, cloud. " * 8
                                + "Requirements: currently studying Computer Science. "
                                + "Applications close 20 October 2030. " + "About us: nice. " * 20)
    audit = make_job(company="BDO", title="Audit Summer Internship", deadline=SOON,
                     description="Responsibilities: " + "audit work for clients, tax and advisory, commercial deals. " * 8
                                 + "Requirements: current students, any degree. " + "About us: nice. " * 20)
    old = make_job(company="Old", title="Software Intern", deadline=PAST)
    jobs = rows(tech, audit, old)
    pipeline._save_previous(jobs)
    excel.write(jobs)                                                  # flag off: `old` is not on the sheet
    calendar_sync.plan_path().write_text('{"actions": []}', encoding="utf-8")
    calendar_sync.ledger_path().write_text(json.dumps({"events": {SOON.isoformat(): {
        "event_id": "ev1", "summary": "2 deadlines", "description": "x", "cleared": False}}}), encoding="utf-8")
    watched = [config.WORKBOOK, config.STATE_DIR / "last_run.json", calendar_sync.ledger_path(),
               calendar_sync.plan_path()]
    before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in watched}

    out = tmp_path / "report"
    s = role_filter.dry_run(out)

    assert {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in watched} == before
    assert sorted(f.name for f in out.iterdir()) == ["assessment.csv", "calendar-impact.csv", "summary.json"]
    assert s["findings"]["role_relevance"] == {"TECHNICAL": 1, "NOT_TECHNICAL": 1, "UNCLEAR": 1}
    assert [line.split()[0] for line in s["sheet"]["leaving"]] == [audit.id]
    assert s["sheet"]["returning_by_reason"] == {"expired - kept, sorted to the bottom with an EXPIRED marker": 1}
    assert "NOT employer-confirmed" in s["coverage"]["rule"]
    day = s["calendar_days_changed"][0]
    assert day["date"] == SOON.isoformat() and "BDO" in day["roles removed"]
    assert s["profile"]["note_for_your_review"]


def test_an_earlier_conflicting_date_gets_its_own_provisional_reminder(on):
    job = technical(rows(make_job(company="Lloyds", title="Software Placement", deadline=SOON + timedelta(days=5)))[0])
    job.deadline_status, job.deadline_provisional = dx.CONFLICT, SOON
    job.deadline_confirmed = job.deadline
    excel.write([job])
    days = calendar_sync.deadline_roles(TODAY)
    assert [r[1] for r in days[job.deadline.isoformat()]] == ["Software Placement"]
    assert days[SOON.isoformat()][0][1].startswith("Software Placement (provisional - the employer's date is")


def test_apply_refuses_until_you_switch_it_on():
    from click.testing import CliRunner

    import main

    before = config.WORKBOOK.exists()
    result = CliRunner().invoke(main.cli, ["assess", "--apply"])
    assert result.exit_code == 2 and "Refused" in result.output
    assert config.WORKBOOK.exists() == before                          # nothing written


def test_applying_keeps_what_the_fetched_employer_pages_showed(on, tmp_path):
    job = rows(make_job(company="Barclays", title="Technology Analyst", deadline=PAST,
                        application_url="https://barclays.wd3.myworkdayjobs.com/x/job/London/Tech_1"))[0]
    job.deadline_status = dx.LISTED
    pages = tmp_path / "pages"
    pages.mkdir()
    later = (TODAY + timedelta(days=30)).isoformat()
    (pages / f"{job.id}.json").write_text(json.dumps({
        "url": job.application_url, "outcome": "checked", "note": "read via the Workday job feed",
        "closed_evidence": "", "checked_at": "2026-09-30T20:00:00",
        "deadline_evidence": [dx.Candidate(date=later, status=dx.OK, source="employer page", url=job.application_url,
                                           method="employer page", employer_page=True, quote="closes").as_dict()]}))
    assert role_filter.merge_fetched([job], pages) == 1
    role_filter.assess_row(job, relevance.Listings([]), TODAY)
    assert job.deadline.isoformat() == later and job.deadline_confirmed.isoformat() == later
    assert job.sheet_decision == role_filter.SHOW and job.check_outcome == "checked"

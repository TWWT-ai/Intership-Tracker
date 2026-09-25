"""Graduation years a posting admits, and the flag for the user's year. No network."""

from openpyxl import load_workbook

from tests.conftest import make_job
from tracker import config, dedupe, excel, graduation, match, verify
from tracker.models import OpenStatus

BLACKROCK_SPRING = ("Our Spring Insight Event is a week-long event taking place in our London and Edinburgh "
                    "offices in April 2027 for students of all academic backgrounds who will graduate with an "
                    "undergraduate or master's degree in 2029.")
BLACKROCK_PLACEMENT = ("BlackRock offers placement opportunities for a duration of 12 months. This programme is "
                       "for candidates graduating in 2029 who are required by their degree to complete a "
                       "long-term internship.")


class TestStated:
    def test_a_stated_year(self):
        years, evidence = graduation.expected_years("2027 Spring Insight Event - EMEA", BLACKROCK_SPRING)
        assert years == [2029] and "2029" in evidence

    def test_a_range_names_every_year_in_it(self):
        text = "Expected graduation between December 2027 - June 2028 and December 2028 - June 2029."
        assert graduation.expected_years("2027 Summer Analyst Program", text)[0] == [2027, 2028, 2029]

    def test_finish_in_counts(self):
        text = ("Currently enrolled in a penultimate year of a Master's or PhD program in a STEM subject, "
                "set to finish in 2028.")
        assert graduation.expected_years("Quant Research Intern - London - 2027", text)[0] == [2028]

    def test_a_year_before_the_graduation_wording_is_not_read(self):
        text = "Our Summer 2027 internship is for students graduating in 2028."
        assert graduation.expected_years("Summer Internship", text)[0] == [2028]

    def test_stated_beats_derived(self):
        text = "Open to first-year students. You must be graduating in 2030."
        assert graduation.expected_years("2027 Spring Week", text)[0] == [2030]

    def test_a_negated_sentence_is_not_read(self):
        text = "This programme is not open to students graduating in 2027."
        assert graduation.expected_years("2027 Spring Week", text) == ([], "")


class TestDerived:
    def test_first_years_on_a_2027_spring_week_graduate_2029_or_2030(self):
        years, evidence = graduation.expected_years("2027 Technology Spring Week", "Open to first-year students.")
        assert years == [2029, 2030] and "2027" in evidence

    def test_a_2027_placement_for_year_in_industry_students_graduates_2029(self):
        text = ("You are an undergraduate studying for a degree in which a year in industry is part of your "
                "structured course syllabus.")
        title = "2027 ISG Central Risk Industrial Placement Internship (London)"
        assert graduation.expected_years(title, text)[0] == [2029]

    def test_penultimate_years_on_a_2027_summer_internship_graduate_2028(self):
        text = "You must be in the penultimate year of your degree."
        assert graduation.expected_years("2027 Technology Summer Internship", text)[0] == [2028]

    def test_no_programme_year_means_no_derived_year(self):
        assert graduation.expected_years("Technology Spring Week", "Open to first-year students.") == ([], "")


class TestFlag:
    def test_your_year_is_flagged_without_retyping_the_row(self):
        """Type comes from what the posting says (tracker/eligibility.py)."""
        job = make_job(title="2027 Summer Internship Programme", work_type="Summer Internship")
        assert graduation.apply(job, BLACKROCK_SPRING, {"graduation_year": 2029})
        assert (job.work_type, job.expected_graduation) == ("Summer Internship", "2029")

    def test_a_placement_for_your_year_stays_a_placement(self):
        job = make_job(title="2027 Placement Programme", work_type="Industrial Placement")
        assert graduation.apply(job, BLACKROCK_PLACEMENT, {"graduation_year": 2029})
        assert job.work_type == "Industrial Placement"

    def test_another_year_changes_only_the_column(self):
        job = make_job(title="2027 Summer Internship", work_type="Summer Internship")
        assert not graduation.apply(job, "For students graduating in 2028.", {"graduation_year": 2029})
        assert (job.work_type, job.expected_graduation) == ("Summer Internship", "2028")

    def test_no_graduation_year_in_the_profile_flags_nothing(self):
        job = make_job(title="2027 Summer Internship", work_type="Summer Internship")
        assert not graduation.apply(job, BLACKROCK_SPRING, {})
        assert (job.work_type, job.expected_graduation) == ("Summer Internship", "2029")

    def test_eligibility_reads_the_stated_years(self):
        job = make_job(title="2027 Summer Internship", work_type="Summer Internship")
        job.expected_graduation = "2028"
        eligible, reasons = match.check_eligibility(job, {"graduation_year": 2029}, config.load("search_filters"))
        assert not eligible and "2028" in reasons[0]


class TestWorkbook:
    def test_expected_graduation_sits_after_spring_week(self):
        labels = [label for label, _ in excel.COLUMNS]
        assert labels[labels.index("Spring Week") + 1] == "Expected Graduation"
        layout = config.load("workbook")["columns"]
        assert layout[layout.index("Spring Week") + 1] == "Expected Graduation"

    def test_the_column_shows_the_years(self):
        job = make_job()
        job.expected_graduation = "2029"
        dedupe.assign_ids([job])
        ws = load_workbook(excel.write([job]))[excel.SHEET_JOBS]
        assert ws.cell(2, excel.col_index()["Expected Graduation"]).value == "2029"

    def test_merging_keeps_the_years(self):
        link = "https://boards.greenhouse.io/acme/jobs/4001111"
        a = make_job(source="trackr", external_id="t1", application_url=link)
        b = make_job(source="careers", external_id="greenhouse:acme:4001111", application_url=link)
        b.expected_graduation = "2029"
        merged, _ = dedupe.deduplicate([a, b])
        assert merged[0].expected_graduation == "2029"


class Page:
    def __init__(self, text, status=200):
        self.text, self.status = text, status

    @property
    def ok(self):
        return 200 <= self.status < 300


class Pages:
    def __init__(self, pages):
        self.pages, self.requested = pages, []
        self.http = self

    def get(self, url, **kwargs):
        self.requested.append(url)
        return self.pages.get(url, Page("", 404))


class TestSweep:
    def test_reads_the_employer_page_only_for_rows_no_feed_described(self):
        job = make_job(title="2027 Spring Insight Programme", work_type="Summer Internship",
                       application_url="https://employer.example.com/job/12345-spring-insight")
        feed_row = make_job(source="careers", application_url="https://employer.example.com/job/2")
        search_page = make_job(application_url="https://employer.example.com/jobs?keyword=intern")
        aggregator = make_job(application_url="https://app.the-trackr.com/uk-tech/summer-internships")
        closed = make_job(application_url="https://employer.example.com/job/3", open_status=OpenStatus.CLOSED)
        session = Pages({job.application_url: Page(f"<html><main><p>{BLACKROCK_SPRING}</p></main></html>")})

        summary = verify.sweep_eligibility(session, [job, feed_row, search_page, aggregator, closed],
                                           config.load("verification"), {"graduation_year": 2029})

        assert (job.work_type, job.expected_graduation) == ("Spring Week", "2029")
        assert session.requested == [job.application_url]
        assert summary.startswith("1 employer pages read")

"""The email status sync: forward only, explicit evidence only, never a guess.

Offline. The Gmail half is Claude's; these tests hand the deterministic half the
evidence file it would have written, and check every rule that keeps a wrong
status out of the workbook.
"""

import json
import shutil
from datetime import date

import pytest
from openpyxl import load_workbook

from tests.conftest import make_job
from tracker import config, excel, pipeline, status_sync
from tracker.models import ApplicationStatus as S


# ---------------------------------------------------------------- helpers

def rows(*jobs):
    """Put jobs in the tracker the way a scan would: last_run.json + workbook."""
    for n, job in enumerate(jobs, start=1):
        job.id = f"INT-{n:06d}"
    pipeline._save_previous(list(jobs))
    excel.write(list(jobs))
    return jobs


def _find(job_id):
    wb = load_workbook(config.WORKBOOK)
    ws = wb[excel.SHEET_JOBS]
    header = [c.value for c in ws[1]]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, header.index("ID") + 1).value == job_id:
            return wb, ws, r, header
    raise AssertionError(f"{job_id} not in sheet")


def cell(job_id, label="Application Status"):
    _, ws, r, header = _find(job_id)
    return ws.cell(r, header.index(label) + 1).value


def set_cell(job_id, value):
    """What the user does in Excel."""
    wb, ws, r, header = _find(job_id)
    ws.cell(r, header.index("Application Status") + 1).value = value
    wb.save(config.WORKBOOK)


def item(company, role, stage, day, mid, year="", domain="hackerrank.com", subject=None):
    return {"company": company, "role_text": role, "stage": stage, "email_date": f"2026-09-{day:02d}",
            "sender_domain": domain, "subject": subject or f"{company} {stage}", "thread_id": f"t-{mid}",
            "message_id": mid, "programme_year": year}


def evidence(*items, ignored=None):
    status_sync.evidence_path().write_text(json.dumps(
        {"emails_checked": len(items) + sum((ignored or {}).values()), "ignored": ignored or {},
         "items": list(items)}), encoding="utf-8")


def sync():
    return status_sync.run()


def barclays():
    return rows(make_job(company="Barclays", title="2027 Technology Summer Analyst"))[0]


# ---------------------------------------------------------------- forward only

def test_forward_progression_is_written_with_its_evidence():
    job = barclays()
    set_cell(job.id, "Application Submitted")
    evidence(item("Barclays", "Technology Summer Analyst", "Online Assessment", 25, "m1"))

    result = sync()

    assert result.wrote and cell(job.id) == "Online Assessment"
    assert cell(job.id, "Status Evidence") == "2026-09-25 hackerrank.com - Barclays Online Assessment"
    assert "Application Submitted -> Online Assessment" in result.text


def test_a_backwards_move_is_refused():
    job = barclays()
    set_cell(job.id, "Video Interview")
    evidence(item("Barclays", "Technology Summer Analyst", "Online Assessment", 25, "m1"))

    result = sync()

    assert cell(job.id) == "Video Interview"
    assert [o.kind for o in result.outcomes] == ["up_to_date"]


@pytest.mark.parametrize("final", ["Offer", "Rejected", "Withdrawn"])
@pytest.mark.parametrize("stage", ["Assessment Centre", "Offer", "Rejected"])
def test_terminal_statuses_are_never_changed(final, stage):
    job = barclays()
    set_cell(job.id, final)
    evidence(item("Barclays", "Technology Summer Analyst", stage, 25, "m1"))

    sync()

    assert cell(job.id) == final


def test_rejection_is_reachable_from_any_open_stage():
    job = barclays()
    set_cell(job.id, "Online Assessment")
    evidence(item("Barclays", "Technology Summer Analyst", "Rejected", 25, "m1", domain="barclays.com"))

    sync()

    assert cell(job.id) == "Rejected"


def test_emails_apply_in_date_order_so_an_older_one_never_wins():
    job = barclays()
    set_cell(job.id, "Application Submitted")
    # Listed out of order on purpose: the sync sorts by email date.
    evidence(item("Barclays", "Technology Summer Analyst", "Video Interview", 26, "m2", domain="hirevue.com"),
             item("Barclays", "Technology Summer Analyst", "Application Submitted", 27, "m3",
                  domain="myworkday.com"),
             item("Barclays", "Technology Summer Analyst", "Online Assessment", 25, "m1"))

    sync()

    assert cell(job.id) == "Video Interview"
    assert cell(job.id, "Status Evidence").startswith("2026-09-26 hirevue.com")


def test_an_older_email_in_a_later_run_cannot_replace_the_status_or_its_evidence():
    job = barclays()
    set_cell(job.id, "Application Submitted")
    evidence(item("Barclays", "Technology Summer Analyst", "Online Assessment", 25, "m1"))
    sync()

    # Arrives late (a delayed or re-surfaced thread) but was sent before the OA.
    evidence(item("Barclays", "Technology Summer Analyst", "Rejected", 20, "m0", domain="barclays.com"))
    result = sync()

    assert cell(job.id) == "Online Assessment"
    assert cell(job.id, "Status Evidence").startswith("2026-09-25 hackerrank.com")
    assert [o.kind for o in result.outcomes] == ["stale"]


def test_status_evidence_changes_only_when_the_status_advances():
    job = barclays()
    set_cell(job.id, "Application Submitted")
    evidence(item("Barclays", "Technology Summer Analyst", "Online Assessment", 25, "m1"))
    sync()
    # Newer email, lower stage: a duplicate confirmation. Nothing moves.
    evidence(item("Barclays", "Technology Summer Analyst", "Application Submitted", 26, "m2",
                  domain="myworkday.com"))
    sync()

    assert cell(job.id, "Status Evidence").startswith("2026-09-25 hackerrank.com")


# ---------------------------------------------------------------- matching

def test_a_different_recruitment_year_goes_to_review():
    job = barclays()
    set_cell(job.id, "Application Submitted")
    evidence(item("Barclays", "Technology Summer Analyst", "Rejected", 25, "m1", year="2026",
                  domain="barclays.com"))

    result = sync()

    assert cell(job.id) == "Application Submitted"
    assert result.outcomes[0].kind == "review" and "2026 round" in result.outcomes[0].reason


def test_two_possible_rows_go_to_review_and_nothing_is_written():
    a, b = rows(make_job(company="Amazon", title="Software Development Engineer Intern", location="London"),
                make_job(company="Amazon", title="Software Development Engineer Intern", location="Edinburgh"))
    set_cell(a.id, "Application Submitted")
    set_cell(b.id, "Application Submitted")               # applied to both: the email names neither city
    evidence(item("Amazon", "Software Development Engineer Intern", "Online Assessment", 25, "m1"))

    result = sync()

    assert cell(a.id) == cell(b.id) == "Application Submitted"
    assert result.outcomes[0].kind == "review" and "could match 2" in result.outcomes[0].reason
    review = load_workbook(config.WORKBOOK)[excel.SHEET_STATUS_REVIEW]
    assert review.max_row == 2 and "Amazon" in (review.cell(2, 6).value or "")


def test_a_role_too_vague_to_name_one_application_goes_to_review():
    a, b = rows(make_job(company="Amazon", title="Software Engineer Intern"),
                make_job(company="Amazon", title="AI/ML Intern"))
    evidence(item("Amazon", "Intern", "Online Assessment", 25, "m1"))

    result = sync()

    assert cell(a.id) is None and cell(b.id) is None
    assert result.outcomes[0].kind == "review"


def test_the_one_open_application_breaks_a_tie_between_listings_of_one_role():
    london, manchester = rows(make_job(company="Lloyds Banking Group", title="Software Engineer Placement",
                                       location="London"),
                              make_job(company="Lloyds Banking Group", title="Software Engineer Placement",
                                       location="Manchester"))
    set_cell(manchester.id, "Application Submitted")
    evidence(item("Lloyds Banking Group", "Software Engineer Placement", "Online Assessment", 25, "m1"))

    sync()

    assert cell(manchester.id) == "Online Assessment" and cell(london.id) is None


def test_an_unknown_company_changes_nothing():
    job = barclays()
    evidence(item("Acme Robotics", "Software Intern", "Online Assessment", 25, "m1"))

    result = sync()

    assert cell(job.id) is None and result.outcomes[0].kind == "no_match"
    assert "No matching application: 1" in result.text


def test_an_unknown_role_changes_nothing():
    job = barclays()
    set_cell(job.id, "Application Submitted")
    evidence(item("Barclays", "Wealth Management Graduate Scheme", "Online Assessment", 25, "m1"))

    result = sync()

    assert cell(job.id) == "Application Submitted" and result.outcomes[0].kind == "review"


def test_an_unclear_stage_goes_to_review_not_a_guess():
    job = barclays()
    set_cell(job.id, "Application Submitted")
    evidence(item("Barclays", "Technology Summer Analyst", "unclear", 25, "m1", domain="barclays.com"))

    result = sync()

    assert cell(job.id) == "Application Submitted"
    assert result.outcomes[0].kind == "review" and "Needs review: 1" in result.text


def test_irrelevant_email_is_counted_but_never_reviewed():
    job = barclays()
    evidence(ignored={"newsletter": 3, "marketing": 2})

    result = sync()

    assert not result.wrote and cell(job.id) is None
    assert "Ignored (no status evidence): marketing 2, newsletter 3" in result.text
    assert "No status changes." in result.text
    assert excel.SHEET_STATUS_REVIEW not in load_workbook(config.WORKBOOK).sheetnames


# ---------------------------------------------------------------- manual edits

def test_a_status_you_edit_after_the_sync_locks_the_row():
    job = barclays()
    set_cell(job.id, "Application Submitted")
    evidence(item("Barclays", "Technology Summer Analyst", "Online Assessment", 25, "m1"))
    sync()

    set_cell(job.id, "Second Stage")                     # you, in Excel
    evidence(item("Barclays", "Technology Summer Analyst", "Assessment Centre", 28, "m2",
                  domain="barclays.com"))
    result = sync()

    assert cell(job.id) == "Second Stage"
    assert cell(job.id, "Status Evidence") == "locked - status edited by you"
    assert [o.kind for o in result.outcomes] == ["locked"]


def test_unlock_hands_the_row_back_from_where_you_left_it():
    job = barclays()
    set_cell(job.id, "Application Submitted")
    evidence(item("Barclays", "Technology Summer Analyst", "Online Assessment", 25, "m1"))
    sync()
    set_cell(job.id, "Second Stage")
    evidence(item("Barclays", "Technology Summer Analyst", "Assessment Centre", 28, "m2",
                  domain="barclays.com"))
    sync()                                               # locks

    status_sync.unlock_row(job.id)
    evidence(item("Barclays", "Technology Summer Analyst", "Assessment Centre", 29, "m3",
                  domain="barclays.com"))
    sync()

    assert cell(job.id) == "Assessment Centre"


def test_unrecognised_text_in_the_cell_is_never_overwritten():
    job = barclays()
    set_cell(job.id, "waiting to hear back")
    evidence(item("Barclays", "Technology Summer Analyst", "Online Assessment", 25, "m1"))

    sync()

    assert cell(job.id) == "waiting to hear back"


# ---------------------------------------------------------------- idempotency and state

def test_the_same_evidence_twice_changes_nothing_the_second_time():
    job = barclays()
    set_cell(job.id, "Application Submitted")
    evidence(item("Barclays", "Technology Summer Analyst", "Online Assessment", 25, "m1"))

    first, second = sync(), sync()

    assert first.wrote and "Applications updated: 1" in first.text
    assert not second.wrote and "Applications updated: 0" in second.text
    assert cell(job.id) == "Online Assessment"


def test_the_ledger_round_trips():
    ledger = {"rows": {"INT-000001": {"status": "Online Assessment", "locked": False,
                                      "evidence": {"email_date": "2026-09-25", "message_id": "m1"}}},
              "processed": {"m1": {"outcome": "changed"}}, "review": {}}
    status_sync.save_ledger(ledger)
    assert status_sync.load_ledger() == ledger


def test_a_dry_run_writes_nothing():
    job = barclays()
    set_cell(job.id, "Application Submitted")
    evidence(item("Barclays", "Technology Summer Analyst", "Online Assessment", 25, "m1"))

    result = status_sync.run(dry_run=True)

    assert "DRY RUN" in result.text and "Proposed changes:" in result.text
    assert cell(job.id) == "Application Submitted"
    assert not status_sync.ledger_path().exists()


# ---------------------------------------------------------------- the workbook

def test_status_updates_is_the_only_way_a_status_is_written():
    job = barclays()
    set_cell(job.id, "Application Submitted")

    excel.write(_same_id(job))
    assert cell(job.id) == "Application Submitted"            # None: exactly as the user left it

    excel.write(_same_id(job), status_updates={job.id: S.OA})
    assert cell(job.id) == "Online Assessment"


def _same_id(job):
    again = make_job(company="Barclays", title="2027 Technology Summer Analyst")
    again.id = job.id
    return [again]


def test_a_synced_status_and_its_evidence_survive_the_nightly_rebuild():
    job = barclays()
    set_cell(job.id, "Application Submitted")
    evidence(item("Barclays", "Technology Summer Analyst", "Online Assessment", 25, "m1"))
    sync()

    excel.write(_same_id(job))                                  # what the 01:00 scan does

    assert cell(job.id) == "Online Assessment"
    assert cell(job.id, "Status Evidence").startswith("2026-09-25 hackerrank.com")


# ---------------------------------------------------------------- guards

def test_an_open_workbook_is_skipped_and_nothing_written():
    job = barclays()
    set_cell(job.id, "Application Submitted")
    evidence(item("Barclays", "Technology Summer Analyst", "Online Assessment", 25, "m1"))
    (config.WORKBOOK.parent / f"~${config.WORKBOOK.name}").write_text("lock")

    result = sync()

    assert result.skipped == "workbook busy" and "skipped - workbook busy" in result.text
    assert cell(job.id) == "Application Submitted"


def test_a_held_sync_lock_is_skipped_and_nothing_written():
    job = barclays()
    set_cell(job.id, "Application Submitted")
    evidence(item("Barclays", "Technology Summer Analyst", "Online Assessment", 25, "m1"))
    status_sync.lock_path().write_text("running")

    result = sync()

    assert result.skipped == "workbook busy" and cell(job.id) == "Application Submitted"


def test_a_failed_backup_writes_nothing(monkeypatch):
    job = barclays()
    set_cell(job.id, "Application Submitted")
    evidence(item("Barclays", "Technology Summer Analyst", "Online Assessment", 25, "m1"))

    def disk_full(*args, **kwargs):
        raise OSError("disk full")
    monkeypatch.setattr(shutil, "copy2", disk_full)

    result = sync()

    assert result.skipped.startswith("backup failed")
    assert cell(job.id) == "Application Submitted"
    assert not status_sync.ledger_path().exists()


def test_a_real_write_is_backed_up_first():
    job = barclays()
    set_cell(job.id, "Application Submitted")
    evidence(item("Barclays", "Technology Summer Analyst", "Online Assessment", 25, "m1"))

    sync()

    backups = list(status_sync.backup_dir().glob("status-sync-*.xlsx"))
    assert len(backups) == 1
    before = load_workbook(backups[0])[excel.SHEET_JOBS]
    header = [c.value for c in before[1]]
    assert before.cell(2, header.index("Application Status") + 1).value == "Application Submitted"


def test_every_run_is_logged_for_the_morning():
    barclays()
    evidence(ignored={"newsletter": 1})

    status_sync.run(when=status_sync.datetime(2026, 9, 25, 23, 0))

    log = (config.LOG_DIR / "status-sync-2026-09-25.txt").read_text(encoding="utf-8")
    assert log.startswith("Application Status Sync - 23:00")


# ---------------------------------------------------------------- telling one round or listing from another

def test_last_years_round_cannot_take_this_years_confirmation():
    old, new = rows(make_job(company="Stripe", title="Software Engineer, Intern", deadline=date(2026, 2, 14)),
                    make_job(company="Stripe", title="Software Engineer, Intern"))
    evidence(item("Stripe", "Software Engineer, Intern", "Application Submitted", 8, "m1", domain="greenhouse-mail.io"))

    sync()

    assert cell(new.id) == "Application Submitted"
    assert old.id not in status_sync.load_ledger()["rows"]     # expired: off the sheet, and untouched


def test_the_listing_you_marked_is_the_application():
    a, b = rows(make_job(company="Palantir", title="Software Engineer, Internship"),
                make_job(company="Palantir", title="Software Engineer, Internship", location="London"))
    set_cell(b.id, "Application Submitted")
    evidence(item("Palantir", "Software Engineer, Internship", "Online Assessment", 20, "m1"))

    sync()

    assert cell(b.id) == "Online Assessment" and cell(a.id) is None


def test_the_exactly_named_programme_beats_its_normalized_twin():
    placement, summer = rows(
        make_job(company="BlackRock", title="2027 Placement Programme - Technology - Software Engineering"),
        make_job(company="BlackRock", title="2027 Summer Internship Programme - Technology - Software Engineering"))
    evidence(item("BlackRock", "2027 Summer Internship Programme - Technology - Software Engineering",
                  "Video Interview", 4, "m1", domain="mail.modernhire.com"))

    sync()

    assert cell(summer.id) == "Video Interview" and cell(placement.id) is None


def test_the_row_the_email_names_exactly_is_chosen_place_aside():
    named, other = rows(make_job(company="Marshall Wace", title="Technology Intern - 2027"),
                        make_job(company="Marshall Wace", title="Technology Intern - Summer 2027"))
    evidence(item("Marshall Wace", "Technology Intern - London - 2027", "Rejected", 24, "m1", year="2027"))

    sync()

    assert cell(named.id) == "Rejected" and cell(other.id) is None


def test_the_one_row_in_the_named_place_is_chosen():
    anywhere, manchester = rows(
        make_job(company="Lloyds Banking Group", title="Data Science and AI Industrial Placement",
                 location="United Kingdom"),
        make_job(company="Lloyds Banking Group", title="Data Science & AI Industrial Placement - Manchester",
                 location="Manchester"))
    evidence(item("Lloyds Banking Group", "Data Science and AI Industrial Placement (Manchester)",
                  "Online Assessment", 22, "m1", domain="lloydsbanking.com"))

    sync()

    assert cell(manchester.id) == "Online Assessment" and cell(anywhere.id) is None


def test_two_equally_plausible_rows_still_go_to_review():
    a, b = rows(make_job(company="Samsara", title="Software Engineering Intern"),
                make_job(company="Samsara", title="Software Engineering Internship"))
    evidence(item("Samsara", "Software Engineering Internship - Summer", "Application Submitted", 20, "m1",
                  domain="greenhouse-mail.io"))

    result = sync()

    assert cell(a.id) is None and cell(b.id) is None and result.outcomes[0].kind == "review"


def test_a_summer_internship_email_never_lands_on_a_placement():
    placement = rows(make_job(company="Bending Spoons", title="Software engineer, industry placement"))[0]
    evidence(item("Bending Spoons", "Software engineer, summer intern UK", "Application Submitted", 31, "m1",
                  domain="bendingspoons.com"))

    result = sync()

    assert cell(placement.id) is None and result.outcomes[0].kind == "review"


def test_shared_boilerplate_is_not_the_same_role():
    product = rows(make_job(company="PIMCO", title="2027 Summer Intern - Product Analyst, EMEA"))[0]
    evidence(item("PIMCO", "2027 Summer Intern - Technology Analyst, Software Engineering, EMEA",
                  "Application Submitted", 18, "m1", year="2027", domain="myworkday.com"))

    result = sync()

    assert cell(product.id) is None and result.outcomes[0].kind == "review"

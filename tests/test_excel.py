"""The workbook: status preservation, colours, one row per job, idempotency."""

from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from tests.conftest import make_job
from tracker import config, dedupe, excel
from tracker.models import ApplicationStatus, OpenStatus, Priority


def cell(ws, row: int, column: str):
    return ws.cell(row=row, column=excel.col_index()[column])


def header(ws) -> list[str]:
    return [c.value for c in ws[1]]


def rows_by_id(path) -> dict[str, int]:
    ws = load_workbook(path)[excel.SHEET_JOBS]
    return {
        ws.cell(r, 1).value: r
        for r in range(2, ws.max_row + 1)
        if ws.cell(r, 1).value
    }


def build(jobs):
    dedupe.assign_ids(jobs)
    return excel.write(jobs)


class TestStructure:
    def test_creates_every_sheet(self):
        path = build([make_job()])
        assert set(load_workbook(path).sheetnames) >= {
            excel.SHEET_JOBS, excel.SHEET_DEADLINES,
            excel.SHEET_RUNS, excel.SHEET_GUIDE, excel.SHEET_LISTS,
        }

    def test_applications_live_on_the_internships_sheet(self):
        """No separate Applications sheet: the status column is the tracker."""
        assert "Applications" not in load_workbook(build([make_job()])).sheetnames

    def test_header_is_frozen_and_filtered(self):
        ws = load_workbook(build([make_job()]))[excel.SHEET_JOBS]
        assert ws.freeze_panes == "C2"
        assert ws.auto_filter.ref is not None

    def test_lists_sheet_is_hidden(self):
        wb = load_workbook(build([make_job()]))
        assert wb[excel.SHEET_LISTS].sheet_state == "hidden"

    def test_application_url_is_a_real_hyperlink(self):
        """The whole point of the dashboard: click through and apply."""
        job = make_job(application_url="https://employer.example.com/apply")
        ws = load_workbook(build([job]))[excel.SHEET_JOBS]
        assert cell(ws, 2, "Application URL").hyperlink.target == "https://employer.example.com/apply"

    def test_formula_injection_is_defused_in_the_sheet(self):
        job = make_job()
        job.notes = "=HYPERLINK(\"http://evil\",\"click\")"
        ws = load_workbook(build([job]))[excel.SHEET_JOBS]
        assert cell(ws, 2, "Notes").value.startswith("'=")


class TestPlatformColours:
    @pytest.mark.parametrize(
        "source,fill",
        [("trackr", "DCE9F8"), ("gradcracker", "FCE6D5"), ("brightnetwork", "DCEFE2")],
    )
    def test_primary_source_cell_is_colour_coded(self, source, fill):
        """Blue for Trackr, orange for Gradcracker, green for Bright Network."""
        ws = load_workbook(build([make_job(source=source)]))[excel.SHEET_JOBS]
        assert fill in cell(ws, 2, "Primary Source").fill.fgColor.rgb

    def test_display_names_are_used_not_slugs(self):
        ws = load_workbook(build([make_job(source="brightnetwork")]))[excel.SHEET_JOBS]
        assert cell(ws, 2, "Primary Source").value == "Bright Network"

    def test_multi_platform_row_is_bold_and_names_every_platform(self):
        deadline = date(2027, 1, 10)
        jobs = [
            make_job(company="NVIDIA Corporation", source="trackr", deadline=deadline),
            make_job(company="Nvidia Ltd", title="SWE Intern", source="gradcracker",
                     deadline=deadline),
        ]
        merged, _ = dedupe.deduplicate(jobs)
        ws = load_workbook(build(merged))[excel.SHEET_JOBS]

        sources = cell(ws, 2, "Sources")
        assert sources.font.bold is True
        assert "Trackr" in sources.value and "Gradcracker" in sources.value

    def test_single_platform_row_is_not_bold(self):
        ws = load_workbook(build([make_job(source="trackr")]))[excel.SHEET_JOBS]
        assert not cell(ws, 2, "Sources").font.bold


class TestOneRowPerJob:
    def test_three_platforms_produce_one_row_with_three_sources(self):
        """Spec 26, end to end: through dedupe and into the sheet."""
        deadline = date(2027, 1, 10)
        jobs = [
            make_job(company="NVIDIA Corporation", title="Software Engineering Intern",
                     source="trackr", deadline=deadline),
            make_job(company="Nvidia Ltd", title="SWE Intern - Summer 2027",
                     source="gradcracker", deadline=deadline),
            make_job(company="NVIDIA", title="Software Engineer Internship",
                     source="brightnetwork", deadline=deadline),
        ]
        merged, _ = dedupe.deduplicate(jobs)
        path = build(merged)

        ws = load_workbook(path)[excel.SHEET_JOBS]
        assert ws.max_row == 2                     # header + exactly one job

        sources = cell(ws, 2, "Sources").value
        assert all(name in sources for name in ("Trackr", "Gradcracker", "Bright Network"))


class TestStatusPreservation:
    def test_manual_status_survives_a_resync(self):
        """Spec 26: a submitted application must still say so after the daily sync."""
        job = make_job()
        path = build([job])

        wb = load_workbook(path)
        ws = wb[excel.SHEET_JOBS]
        cell(ws, 2, "Application Status").value = "Application Submitted"
        cell(ws, 2, "Notes").value = "Applied via referral"
        wb.save(path)

        # The next sync re-finds the job; adapters always emit FOUND.
        again = make_job()
        again.id = job.id
        assert again.status is ApplicationStatus.FOUND
        excel.write([again])

        ws = load_workbook(path)[excel.SHEET_JOBS]
        assert cell(ws, 2, "Application Status").value == "Application Submitted"
        assert cell(ws, 2, "Notes").value == "Applied via referral"

    def test_status_survives_the_job_appearing_on_a_new_platform(self):
        job = make_job(source="trackr", deadline=date(2027, 1, 10))
        path = build([job])

        wb = load_workbook(path)
        cell(wb[excel.SHEET_JOBS], 2, "Application Status").value = "Application Submitted"
        wb.save(path)

        merged, _ = dedupe.deduplicate([
            make_job(source="trackr", deadline=date(2027, 1, 10)),
            make_job(source="gradcracker", title="SWE Intern", deadline=date(2027, 1, 10)),
        ])
        dedupe.assign_ids(merged)
        excel.write(merged)

        ws = load_workbook(path)[excel.SHEET_JOBS]
        assert cell(ws, 2, "Application Status").value == "Application Submitted"

    def test_applied_row_is_demoted_from_high_priority(self):
        """An applied job must not keep occupying the top of the sheet."""
        job = make_job(deadline=date.today() + timedelta(days=2),
                       open_status=OpenStatus.OPEN)
        job.priority = Priority.HIGH
        path = build([job])

        wb = load_workbook(path)
        cell(wb[excel.SHEET_JOBS], 2, "Application Status").value = "Application Submitted"
        wb.save(path)

        again = make_job(deadline=date.today() + timedelta(days=2), open_status=OpenStatus.OPEN)
        again.id = job.id
        again.priority = Priority.HIGH
        excel.write([again])

        ws = load_workbook(path)[excel.SHEET_JOBS]
        assert cell(ws, 2, "Priority").value == "LOW"

    def test_an_older_applications_sheet_is_left_alone(self):
        """Workbooks from before the merge may still hold one. It is the user's."""
        path = build([make_job()])

        wb = load_workbook(path)
        wb.create_sheet("Applications").append(["APP-1", "INT-000001", "NVIDIA", "SWE Intern"])
        wb.save(path)

        excel.write([make_job()])

        ws = load_workbook(path)["Applications"]
        assert ws.cell(1, 1).value == "APP-1"

    def test_every_later_stage_survives_a_resync(self):
        job = make_job()
        path = build([job])
        for stage in excel.DROPDOWN_STATUSES:
            wb = load_workbook(path)
            cell(wb[excel.SHEET_JOBS], 2, "Application Status").value = str(stage)
            wb.save(path)

            again = make_job()
            again.id = job.id
            excel.write([again])
            ws = load_workbook(path)[excel.SHEET_JOBS]
            assert cell(ws, 2, "Application Status").value == str(stage)

    def test_clearing_the_cell_resets_to_not_applied(self):
        job = make_job()
        job.status = ApplicationStatus.APPLIED
        path = build([job])

        wb = load_workbook(path)
        cell(wb[excel.SHEET_JOBS], 2, "Application Status").value = None
        wb.save(path)

        job.status = ApplicationStatus.APPLIED     # e.g. reloaded from last run
        excel.write([job])
        assert cell(load_workbook(path)[excel.SHEET_JOBS], 2, "Application Status").value is None


class TestApplicationStatusColumn:
    @pytest.mark.parametrize(
        "old,new",
        [("APPLIED", "Application Submitted"), ("OA", "Online Assessment"),
         ("REJECTED", "Rejected"), ("rejected", "Rejected"), ("offer", "Offer")],
    )
    def test_old_codes_are_carried_into_the_new_wording(self, old, new):
        """Statuses typed before the wording changed must not be dropped."""
        job = make_job()
        path = build([job])
        wb = load_workbook(path)
        cell(wb[excel.SHEET_JOBS], 2, "Application Status").value = old
        wb.save(path)

        again = make_job()
        again.id = job.id
        excel.write([again])
        assert cell(load_workbook(path)[excel.SHEET_JOBS], 2, "Application Status").value == new

    def test_a_workbook_with_the_old_status_header_is_read(self):
        job = make_job()
        path = build([job])
        wb = load_workbook(path)
        ws = wb[excel.SHEET_JOBS]
        ws.cell(1, excel.col_index()["Application Status"]).value = "Status"
        cell(ws, 2, "Application Status").value = "OA"
        wb.save(path)

        again = make_job()
        again.id = job.id
        excel.write([again])
        assert cell(load_workbook(path)[excel.SHEET_JOBS], 2, "Application Status").value == "Online Assessment"

    def test_not_applied_is_a_blank_cell(self):
        ws = load_workbook(build([make_job()]))[excel.SHEET_JOBS]
        assert cell(ws, 2, "Application Status").value is None

    def test_sits_just_before_its_evidence_then_notes(self):
        """Application Status, the email behind it, then the user's Notes."""
        labels = [label for label, _ in excel.columns()]
        at = labels.index("Application Status")
        assert labels[at + 1: at + 3] == ["Status Evidence", "Notes"]

    def test_dropdown_offers_every_stage(self):
        wb = load_workbook(build([make_job()]))
        offered = [c.value for c in wb[excel.SHEET_LISTS]["A"][1:]]
        assert offered == ["Application Submitted", "Online Assessment", "Video Interview",
                           "Second Stage", "Assessment Centre", "Offer", "Rejected", "Withdrawn"]

    @pytest.mark.parametrize("work_type", ["Summer Internship", "Industrial Placement", "Spring Week"])
    def test_status_cell_takes_the_work_type_colour(self, work_type):
        job = make_job(work_type=work_type)
        job.status = ApplicationStatus.OA
        ws = load_workbook(build([job]))[excel.SHEET_JOBS]
        assert excel._flag_styles()[work_type]["fill"] in cell(ws, 2, "Application Status").fill.fgColor.rgb

    def test_generic_internship_status_is_not_coloured(self):
        """No specific type means no colour, rather than a guessed one."""
        job = make_job(work_type="Internship")
        job.status = ApplicationStatus.APPLIED
        ws = load_workbook(build([job]))[excel.SHEET_JOBS]
        assert cell(ws, 2, "Application Status").fill.fgColor.rgb in (None, "00000000")

    def test_unrecognised_text_is_not_guessed_into_a_stage(self):
        assert ApplicationStatus.parse("maybe later") is None


class TestIdempotency:
    def test_running_twice_creates_no_duplicate_rows(self):
        """Spec 26: the daily workflow must be safe to run repeatedly."""
        jobs = [make_job(company="NVIDIA"), make_job(company="Arm", title="Robotics Intern")]
        path = build(jobs)
        first = rows_by_id(path)

        build([make_job(company="NVIDIA"), make_job(company="Arm", title="Robotics Intern")])
        second = rows_by_id(path)

        assert len(second) == 2
        assert set(first) == set(second)          # same ids, no renumbering

    def test_ten_consecutive_runs_stay_stable(self):
        for _ in range(10):
            build([make_job()])
        assert len(rows_by_id(config.WORKBOOK)) == 1


class TestSorting:
    def test_high_priority_first_and_closed_last(self):
        low = make_job(company="LowCo", title="Robotics Intern")
        low.priority = Priority.LOW
        low.open_status = OpenStatus.OPEN

        high = make_job(company="HighCo", title="Backend Engineering Intern")
        high.priority = Priority.HIGH
        high.open_status = OpenStatus.OPEN
        high.deadline = date.today() + timedelta(days=3)

        closed = make_job(company="ClosedCo", title="Frontend Engineering Intern")
        closed.open_status = OpenStatus.CLOSED

        ordered = excel.sort_jobs([closed, low, high])
        assert ordered[0].company == "HighCo"
        assert ordered[-1].company == "ClosedCo"

    def test_expired_jobs_sort_below_live_ones(self):
        expired = make_job(company="ExpiredCo", title="Backend Engineering Intern",
                           deadline=date.today() - timedelta(days=3),
                           open_status=OpenStatus.OPEN)
        live = make_job(company="LiveCo", title="Robotics Intern",
                        deadline=date.today() + timedelta(days=30),
                        open_status=OpenStatus.OPEN)
        live.priority = Priority.MEDIUM
        expired.priority = Priority.MEDIUM

        ordered = excel.sort_jobs([expired, live])
        assert ordered[0].company == "LiveCo"

    def test_closed_jobs_are_kept_not_deleted(self):
        jobs = [make_job(company="ClosedCo", open_status=OpenStatus.CLOSED)]
        ws = load_workbook(build(jobs))[excel.SHEET_JOBS]
        assert ws.max_row == 2
        assert cell(ws, 2, "Open Status").value == "CLOSED"


class TestLayout:
    @pytest.fixture
    def layout(self, monkeypatch):
        """Set the configured column list for one test."""
        def use(labels):
            real = excel._layout()
            monkeypatch.setattr(excel, "_layout", lambda: {**real, "columns": labels})
        return use

    def test_only_configured_columns_are_written(self, layout):
        """A column the user removed must not come back on the next run."""
        layout(["ID", "Company", "Job Title", "Application Status", "Notes"])
        ws = load_workbook(build([make_job()]))[excel.SHEET_JOBS]
        assert header(ws) == ["ID", "Company", "Job Title", "Application Status", "Notes"]

    def test_id_and_status_are_always_kept(self, layout):
        """Status is matched back by ID; dropping either would lose it."""
        layout(["Company", "Job Title"])
        ws = load_workbook(build([make_job()]))[excel.SHEET_JOBS]
        assert header(ws) == ["ID", "Company", "Job Title", "Application Status"]

    def test_a_forced_status_column_goes_before_notes(self, layout):
        layout(["Company", "Notes"])
        ws = load_workbook(build([make_job()]))[excel.SHEET_JOBS]
        assert header(ws) == ["ID", "Company", "Application Status", "Notes"]

    def test_unknown_columns_are_ignored(self, layout):
        layout(["ID", "Company", "No Such Column", "Application Status"])
        ws = load_workbook(build([make_job()]))[excel.SHEET_JOBS]
        assert header(ws) == ["ID", "Company", "Application Status"]

    def test_status_survives_a_resync_with_a_trimmed_layout(self, layout):
        layout(["ID", "Company", "Application Status", "Notes"])
        job = make_job()
        path = build([job])

        wb = load_workbook(path)
        cell(wb[excel.SHEET_JOBS], 2, "Application Status").value = "Application Submitted"
        cell(wb[excel.SHEET_JOBS], 2, "Notes").value = "Referral from Sam"
        wb.save(path)

        again = make_job()
        again.id = job.id
        excel.write([again])

        ws = load_workbook(path)[excel.SHEET_JOBS]
        assert cell(ws, 2, "Application Status").value == "Application Submitted"
        assert cell(ws, 2, "Notes").value == "Referral from Sam"


class TestFlagColumns:
    @pytest.mark.parametrize(
        "work_type,column",
        [("Summer Internship", "Summer Internship"),
         ("Industrial Placement", "Industrial Placement"),
         ("Spring Week", "Spring Week")],
    )
    def test_work_type_marks_only_its_own_column(self, work_type, column):
        ws = load_workbook(build([make_job(work_type=work_type)]))[excel.SHEET_JOBS]
        for label in ("Summer Internship", "Industrial Placement", "Spring Week"):
            assert cell(ws, 2, label).value == ("Yes" if label == column else None)

    def test_generic_internship_marks_no_type_column(self):
        """'Internship' could be either; picking one would be a guess."""
        ws = load_workbook(build([make_job(work_type="Internship")]))[excel.SHEET_JOBS]
        for label in ("Summer Internship", "Industrial Placement", "Spring Week"):
            assert cell(ws, 2, label).value is None

    def test_each_flag_column_has_its_own_colour(self):
        job = make_job(work_type="Summer Internship")
        job.cv, job.written_answers, job.cover_letter = "Yes", "Yes", "Yes"
        ws = load_workbook(build([job]))[excel.SHEET_JOBS]

        fills = {label: cell(ws, 2, label).fill.fgColor.rgb
                 for label in ("Summer Internship", "CV", "WA", "CL")}
        assert len(set(fills.values())) == 4
        for label, spec in excel._flag_styles().items():
            if label in fills:
                assert spec["fill"] in fills[label]
                assert spec["header"] in cell(ws, 1, label).fill.fgColor.rgb

    def test_materials_are_written_as_stated(self):
        job = make_job()
        job.cv, job.written_answers, job.cover_letter = "Yes", "Optional", "No"
        ws = load_workbook(build([job]))[excel.SHEET_JOBS]

        assert cell(ws, 2, "CV").value == "Yes"
        assert cell(ws, 2, "CV").font.bold
        assert cell(ws, 2, "WA").value == "Optional"
        assert cell(ws, 2, "CL").value == "No"
        assert cell(ws, 2, "CL").fill.fgColor.rgb in (None, "00000000")

    def test_unstated_material_stays_blank(self):
        """Not stated is not 'No'."""
        ws = load_workbook(build([make_job()]))[excel.SHEET_JOBS]
        assert cell(ws, 2, "WA").value is None

    def test_notes_is_left_for_the_user(self):
        ws = load_workbook(build([make_job()]))[excel.SHEET_JOBS]
        assert cell(ws, 2, "Notes").value is None


class TestOpenWorkbookFallback:
    def test_locked_workbook_diverts_instead_of_losing_the_run(self, monkeypatch):
        from openpyxl import Workbook

        real_save = Workbook.save
        calls = {"n": 0}

        def flaky_save(self, path):
            calls["n"] += 1
            if calls["n"] == 1:
                raise PermissionError("file is open in Excel")
            return real_save(self, path)

        monkeypatch.setattr(Workbook, "save", flaky_save)
        written = excel.write([make_job()])
        assert written.name.endswith(".PENDING.xlsx")

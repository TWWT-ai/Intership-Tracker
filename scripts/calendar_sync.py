"""Deadline reminders in Google Calendar: what to create or change tonight.

The Python half never touches Google - it holds no token (rule 5). It reads
the workbook, decides one event per deadline day, and diffs that against the
ledger of events already written. A Claude run with the Calendar connector
carries out the plan, then records what it wrote:

    python scripts/calendar_sync.py plan                    -> data/state/calendar_plan.json
    python scripts/calendar_sync.py record <results.json>   {"YYYY-MM-DD": "<event id>", ...}

Only dates in the results file are recorded, so an action that failed is
simply planned again the next night. Roles the user has already applied to
(any Application Status) are not reminded about. A day whose roles have all
gone is cleared - reminders off, text says so - never deleted.
"""
from __future__ import annotations

import datetime as dt
import json
import sys
from collections import OrderedDict, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from tracker import config  # noqa: E402
from tracker.models import ApplicationStatus  # noqa: E402

#: 7, 3 and 1 days before, in minutes before the 08:30 event.
REMINDER_MINUTES = [7 * 1440, 3 * 1440, 1440]
EVENT_TIME = "08:30"
TIME_ZONE = "Europe/London"
FOOTER = ("From your internship tracker (data/internships.xlsx). "
          "Roles you have already applied to are not listed.")


def ledger_path() -> Path:
    return config.STATE_DIR / "calendar_events.json"


def plan_path() -> Path:
    return config.STATE_DIR / "calendar_plan.json"


def load_ledger() -> dict:
    path = ledger_path()
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return {"events": {}}


def _plain(text) -> str:
    return " ".join(str(text or "").replace("–", "-").replace("—", "-").split())


def workbook_rows() -> list[tuple]:
    """(deadline, status cell, company, title, id) for every Internships row,
    plus a separate "(provisional)" row for each earlier conflicting date the
    Assessment sheet keeps apart from the employer's deadline."""
    wb = load_workbook(config.WORKBOOK, read_only=True)
    rows = wb["Internships"].iter_rows(values_only=True)
    col = {name: i for i, name in enumerate(next(rows)) if name}
    out = [(row[col["Deadline"]], row[col["Application Status"]], row[col["Company"]], row[col["Job Title"]],
            row[col["ID"]]) for row in rows]
    if "Assessment" in wb.sheetnames:
        status = {r[4]: r[1] for r in out}
        rows = wb["Assessment"].iter_rows(values_only=True)
        acol = {name: i for i, name in enumerate(next(rows)) if name}
        for row in rows:
            job_id, provisional = row[acol["ID"]], row[acol["Provisional (kept apart)"]]
            if provisional and job_id in status:
                out.append((provisional, status[job_id], row[acol["Company"]],
                            f"{row[acol['Job Title']]} (provisional - the employer's date is "
                            f"{_plain(row[acol['Deadline']])[:10]})", job_id))
    return out


def deadline_roles(today: dt.date, rows: list[tuple] | None = None) -> dict[str, list[tuple[str, str, str]]]:
    """{YYYY-MM-DD: [(company, title, id)]} for roles still to apply to.

    `rows` defaults to the workbook; the assessment dry run passes the rows
    the sheet would hold, to report reminders a change would touch."""
    days: dict[str, list] = defaultdict(list)
    for deadline, status, company, title, job_id in (workbook_rows() if rows is None else rows):
        if isinstance(deadline, dt.datetime):
            deadline = deadline.date()
        if not isinstance(deadline, dt.date) or deadline < today:
            continue
        # Blank is "not applied"; anything else, including text the dropdown
        # does not know, is the user's own record and needs no reminder.
        if ApplicationStatus.parse(status) is not ApplicationStatus.FOUND:
            continue
        days[deadline.isoformat()].append((_plain(company), _plain(title), str(job_id)))
    return days


def event_text(roles: list[tuple[str, str, str]]) -> tuple[str, str]:
    """(summary, description) for one deadline day."""
    roles = sorted(roles, key=lambda r: (r[0].casefold(), r[1].casefold()))
    companies = list(OrderedDict.fromkeys(r[0] for r in roles))
    if len(roles) == 1:
        summary = f"Deadline: {roles[0][1]} - {roles[0][0]}"[:120]
        head = "Application deadline today:"
    else:
        summary = f"{len(roles)} deadlines: " + ", ".join(companies[:3]) + (
            f" and {len(companies) - 3} more" if len(companies) > 3 else "")
        head = f"{len(roles)} application deadlines today:"
    lines = "\n".join(f"- {title} - {company} ({job_id})" for company, title, job_id in roles)
    return summary, f"{head}\n{lines}\n\n{FOOTER}"


CLEARED_SUMMARY = "No open deadlines left today"
CLEARED_DESCRIPTION = ("Every role that had its deadline today has been applied to or removed "
                       "from the tracker. Reminders for this day are switched off.")


def plan(today: dt.date | None = None, rows: list[tuple] | None = None) -> dict:
    """The actions that bring the calendar in line with the workbook (or with
    `rows`, for a dry run). Reads the ledger; writes nothing."""
    today = today or dt.date.today()
    events = load_ledger()["events"]
    wanted = {day: event_text(roles) for day, roles in deadline_roles(today, rows).items()}
    actions = []
    for day in sorted(set(wanted) | {d for d in events if d >= today.isoformat()}):
        have = events.get(day)
        if day in wanted:
            summary, description = wanted[day]
            if not have:
                kind = "create"
            elif have.get("cleared") or (have.get("summary"), have.get("description")) != (summary, description):
                kind = "update"
            else:
                continue
            actions.append({"date": day, "action": kind, "event_id": (have or {}).get("event_id"),
                            "summary": summary, "description": description, "reminders": REMINDER_MINUTES})
        elif not have.get("cleared"):
            actions.append({"date": day, "action": "clear", "event_id": have["event_id"],
                            "summary": CLEARED_SUMMARY, "description": CLEARED_DESCRIPTION, "reminders": []})
    return {"generated": dt.datetime.now().isoformat(timespec="seconds"), "time": EVENT_TIME,
            "time_zone": TIME_ZONE, "actions": actions}


def record(results: dict[str, str]) -> int:
    """Mark the planned actions carried out (date -> event id) as written."""
    pending = json.loads(plan_path().read_text(encoding="utf-8"))["actions"]
    ledger = load_ledger()
    done = 0
    for action in pending:
        event_id = results.get(action["date"])
        if not event_id:
            continue
        ledger["events"][action["date"]] = {"event_id": event_id, "summary": action["summary"],
                                            "description": action["description"],
                                            "cleared": action["action"] == "clear"}
        done += 1
    ledger["updated"] = dt.datetime.now().isoformat(timespec="seconds")
    ledger_path().write_text(json.dumps(ledger, indent=1, ensure_ascii=False), encoding="utf-8")
    return done


def main(argv: list[str]) -> int:
    if argv[:1] == ["plan"]:
        result = plan()
        plan_path().write_text(json.dumps(result, indent=1, ensure_ascii=False), encoding="utf-8")
        counts = {k: sum(1 for a in result["actions"] if a["action"] == k) for k in ("create", "update", "clear")}
        print(f"Calendar plan: {counts['create']} to create, {counts['update']} to update, "
              f"{counts['clear']} to clear -> {plan_path()}")
        return 0
    if argv[:1] == ["record"] and len(argv) == 2:
        results = json.loads(Path(argv[1]).read_text(encoding="utf-8"))
        print(f"Calendar ledger: {record(results)} event(s) recorded")
        return 0
    print(__doc__)
    return 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))

"""Before/after the graduation run, read-only. Columns found by header label."""
import collections
import json
from pathlib import Path

from openpyxl import load_workbook

SCRATCH = Path(__file__).parent
ROOT = Path(r"D:\School\Personal\Intership tracker")
BEFORE = SCRATCH / "backups" / "internships-before-graduation.xlsx"
AFTER = ROOT / "data" / "internships.xlsx"
TYPES = ("Summer Internship", "Industrial Placement", "Spring Week")
TARGET = "2029"


def ascii(text):
    return str(text if text is not None else "").encode("ascii", "replace").decode()


def rows(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Internships"]
    it = ws.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(it)]
    out = {str(dict(zip(header, v))["ID"]): dict(zip(header, v)) for v in it if v and v[0]}
    wb.close()
    return header, out


def kind(row):
    return [t for t in TYPES if str(row.get(t) or "").strip() not in ("", "No")]


hb, before = rows(BEFORE)
ha, after = rows(AFTER)
print(f"rows: before {len(before)} -> after {len(after)}")
i = ha.index("Spring Week") if "Spring Week" in ha else -1
print("column after Spring Week:", ha[i + 1] if i >= 0 else None, "| columns otherwise unchanged:",
      [h for h in ha if h != "Expected Graduation"] == hb)

lost, altered, count = [], [], 0
for job_id, old in before.items():
    for col in ("Application Status", "Notes"):
        value = str(old.get(col) or "").strip()
        if not value:
            continue
        count += 1
        new = after.get(job_id)
        if new is None:
            lost.append((job_id, col, value))
        elif str(new.get(col) or "").strip() != value:
            altered.append((job_id, col, value, new.get(col)))
print(f"user-set cells before: {count} | row missing after: {len(lost)} | changed: {len(altered)}")
for item in lost + altered:
    print("   !!", ascii(item))

grad = {k: r for k, r in after.items() if str(r.get("Expected Graduation") or "").strip()}
admit = {k: r for k, r in grad.items() if TARGET in str(r.get("Expected Graduation"))}
print(f"\nrows with Expected Graduation: {len(grad)} | admitting {TARGET}: {len(admit)}")
print("years shown:", collections.Counter(str(r.get("Expected Graduation")) for r in grad.values()).most_common(10))
print("type flags before:", collections.Counter(t for r in before.values() for t in kind(r)))
print("type flags after: ", collections.Counter(t for r in after.values() for t in kind(r)))

state = {j["id"]: j for j in json.loads((ROOT / "data/state/last_run.json").read_text(encoding="utf-8"))["jobs"]}
print(f"\nrows admitting {TARGET} (company | title | type | open | deadline | evidence):")
for k, r in sorted(admit.items(), key=lambda kv: (str(kv[1].get("Company")), str(kv[1].get("Job Title")))):
    evidence = (state.get(k) or {}).get("year_of_study") or ""
    print(f"   {k} | {ascii(r.get('Company'))[:22]:22} | {ascii(r.get('Job Title'))[:52]:52} | "
          f"{','.join(kind(r)) or '-':20} | {ascii(r.get('Open Status')):7} | {ascii(r.get('Deadline'))[:10]:10} | "
          f"{ascii(evidence)[:110]}")

new_ids = sorted(set(after) - set(before))
print(f"\nnew rows: {len(new_ids)} | by sources: "
      f"{collections.Counter(ascii(after[i].get('Sources')) for i in new_ids).most_common(8)}")
retyped = [(k, kind(before[k]), kind(after[k])) for k in set(before) & set(after) if kind(before[k]) != kind(after[k])]
print(f"existing rows whose type changed: {len(retyped)}")
for k, b, a in sorted(retyped)[:60]:
    print(f"   {k} | {ascii(after[k].get('Company'))[:22]:22} | {ascii(after[k].get('Job Title'))[:48]:48} | {b} -> {a} | grad {ascii(after[k].get('Expected Graduation'))}")

print("\nBlackRock rows:")
for k, r in sorted(after.items()):
    if r.get("Company") == "BlackRock":
        print(f"   {k} | {ascii(r.get('Job Title'))[:48]:48} | {','.join(kind(r)) or '-':20} | grad {ascii(r.get('Expected Graduation')):12} | "
              f"{ascii(r.get('Open Status')):7} | {ascii(r.get('Deadline'))[:10]} | status {ascii(r.get('Application Status'))}")

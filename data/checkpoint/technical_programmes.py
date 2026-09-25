"""One row per technical programme, company-level cells merged. Formatting + filtering only.

Source: the last one-row-per-programme workbook (the backup taken before the
company consolidation). Any Application Status or Note the user has since typed
into the consolidated workbook is carried back by programme ID.

Non-technical programmes are removed by their Job Title (Skills as a tie-break).
A title naming anything technical is kept even inside a finance division;
unclear titles are kept and listed for review. Company is merged across each
company's rows; Match Score and Location only where every row holds the same
value. The draft is re-read and verified before the workbook is replaced.

    pipenv run python data/checkpoint/technical_programmes.py --check    # verify, write nothing
    pipenv run python data/checkpoint/technical_programmes.py            # replace the workbook
"""
import csv
import datetime as dt
import os
import re
import shutil
import sys
from collections import Counter, OrderedDict
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(r"D:\School\Personal\Intership tracker")
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

from tracker import normalize  # noqa: E402

CURRENT = ROOT / "data" / "internships.xlsx"
SOURCE = ROOT / "data" / "checkpoint" / "backups" / "internships-before-company-consolidation.xlsx"
BACKUP = ROOT / "data" / "checkpoint" / "backups" / "internships-before-technical-filter.xlsx"
if "--after-daily" in sys.argv:
    # The daily scan has just rewritten the workbook one row per programme: that is the source.
    SOURCE = CURRENT
    BACKUP = ROOT / "data" / "checkpoint" / "backups" / "internships-daily-unformatted.xlsx"
DRAFT = ROOT / "data" / "checkpoint" / "internships-technical-draft.xlsx"
REMOVED_CSV = ROOT / "data" / "checkpoint" / "removed-non-technical.csv"
REVIEW_CSV = ROOT / "data" / "checkpoint" / "kept-for-review.csv"
SHEET = "Internships"

COMPANY_LEVEL = ["Company", "Match Score", "Location"]
LINE_HEIGHT, MAX_ROW_HEIGHT = 15, 409

#: Same-employer spellings confirmed earlier today (shared programmes or careers links).
SAME_COMPANY = {"Maven": "Maven Securities", "Marshall Wace Internship Programmes": "Marshall Wace",
                "Bny": "BNY Mellon", "Zurich": "Zurich Insurance", "Cr\u00e9dit Agricole CIB": "Cr\u00e9dit Agricole",
                "Honeywell Aerospace": "Honeywell", "Symphony Communication Services": "Symphony"}
AGGREGATORS = {"Gradcracker", "Bright Network", "Higherin", "TargetJobs"}


def words(*phrases):
    return re.compile(r"\b(?:" + "|".join(re.escape(p).replace(r"\ ", r"[\s&/-]+") for p in phrases) + r")\b",
                      re.IGNORECASE)


# Phrases where a technical-looking word is really a business or HR term.
NOT_TECHNICAL_PHRASES = words("business development", "leadership development", "talent development",
                              "development programme", "development program", "people development",
                              "graduate development", "professional development", "sales engineer",
                              "securities", "non technical")
COMPUTING = words("software", "developer", "developers", "development", "programming", "programmer", "computer",
                  "computing", "technology", "technologies", "tech", "technical", "it", "digital", "data", "analytics",
                  "ai", "artificial intelligence", "machine learning", "ml", "deep learning", "nlp", "cyber",
                  "cybersecurity", "security", "infosec", "cloud", "devops", "sre",
                  "site reliability", "network", "networks", "networking", "systems", "embedded", "firmware",
                  "hardware", "electronic", "electronics", "electrical", "robotics", "automation", "simulation",
                  "algorithmic", "algorithms", "platform", "web", "frontend", "front end", "backend", "back end",
                  "full stack", "mobile", "ios", "android", "fpga", "silicon", "semiconductor", "chip", "controls",
                  "control systems", "autonomy", "autonomous", "computational", "informatics", "quantitative developer",
                  "quant developer", "quantitative technology", "research engineer", "r&d", "blockchain", "sap",
                  "salesforce developer", "test engineer", "qa", "quality assurance", "ux", "ui", "gaming",
                  "game", "games", "vfx", "graphics", "cad", "stem", "engineering technology")
NON_TECHNICAL = words("finance", "financial", "corporate finance", "accounting", "accountancy", "accountant",
                      "treasury", "investment banking", "banking", "wealth", "asset management", "private equity",
                      "equity research", "fixed income", "credit", "markets", "global markets", "sales", "marketing",
                      "marketer", "brand", "communications", "public relations", "human resources", "hr", "people",
                      "talent", "recruitment", "legal", "law", "lawyer", "solicitor", "paralegal", "compliance",
                      "risk", "audit", "assurance", "tax", "insurance", "actuarial", "underwriting", "claims",
                      "operations", "business", "commercial", "strategy", "consulting", "advisory", "consultant",
                      "project management", "programme management", "procurement", "supply chain", "logistics",
                      "customer", "client", "relationship management", "investment", "investments", "economics",
                      "economist", "policy", "property", "real estate", "surveying", "surveyor", "planning",
                      "corporate", "administration", "events", "retail", "merchandising", "buying", "hospitality",
                      "vacation scheme", "training contract", "product management", "product manager", "trading",
                      "trader", "structuring", "portfolio", "intellectual property")
#: Finance, sales and legal functions. A title pairing one of these with only a
#: broad technical word ("Investment Banking: Data Analytics", "Legal Technology")
#: is left for review rather than kept outright.
STRONG_NON_TECHNICAL = words("sales", "relationship management", "credit", "private equity", "investment banking",
                             "portfolio", "equity research", "wealth", "asset management", "insurance", "markets",
                             "global markets", "trading", "structuring", "investment", "investments", "legal",
                             "intellectual property")
#: Words that make a programme technical whatever division it sits in.
STRONG_TECHNICAL = words("software", "developer", "developers", "programming", "programmer", "computer science",
                         "machine learning", "ai", "artificial intelligence", "deep learning", "nlp", "embedded",
                         "firmware", "hardware", "fpga", "quantitative developer", "quant developer", "devops", "sre",
                         "site reliability", "cloud", "robotics", "electronic", "electronics", "electrical")
OTHER_ENGINEERING = words("engineering", "engineer", "engineers", "mechanical", "civil", "structural", "chemical",
                          "process", "manufacturing", "materials", "aerospace", "aeronautical", "naval", "marine",
                          "nuclear", "building services", "mining", "geotechnical", "environmental", "energy", "power",
                          "science", "scientist", "physics", "mathematics", "maths", "design", "performance",
                          "manufacture", "production", "quality", "test", "testing", "safety")
QUANT = words("quant", "quants", "quantitative")
SKILL_SIGNAL = words("python", "java", "c++", "c#", "javascript", "typescript", "sql", "go", "rust", "matlab",
                     "linux", "aws", "azure", "kubernetes", "docker", "react", "pytorch", "tensorflow", "git")


def classify(title: str, skills: str) -> tuple[str, str]:
    """(keep | remove | review, why) from the programme's own title, then its skills."""
    clean = NOT_TECHNICAL_PHRASES.sub(" ", title or "")
    tech = COMPUTING.search(clean)
    strong_tech = STRONG_TECHNICAL.search(clean)
    strong_non = STRONG_NON_TECHNICAL.search(clean)
    quant = QUANT.search(clean)
    if tech:
        if strong_non and not strong_tech:
            return "review", f"technical word ({tech.group(0)}) in a {strong_non.group(0)} programme"
        if quant and not strong_tech:
            return "review", f"quantitative role ({quant.group(0)}) - technical or finance?"
        return "keep", f"technical title: {tech.group(0)}"
    non_tech = NON_TECHNICAL.search(clean)
    engineering = OTHER_ENGINEERING.search(clean)
    if non_tech and not (quant or engineering):
        return "remove", f"non-technical function: {non_tech.group(0)}"
    if SKILL_SIGNAL.search(skills or "") and not non_tech:
        return "keep", "technical skills listed"
    if quant:
        return "review", f"quantitative role ({quant.group(0)}) - technical or trading?"
    if non_tech and engineering:
        return "review", f"mixed: {non_tech.group(0)} + {engineering.group(0)}"
    if engineering:
        return "review", f"non-computing engineering/science: {engineering.group(0)}"
    return "review", "title does not name a function"


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    return str(value).strip()


def company_key(name: str) -> str:
    return normalize.company(name)[1] or name.casefold().strip()


# ---------------------------------------------------------------- read source (one row per programme)
src_wb = load_workbook(SOURCE)
src = src_wb[SHEET]
old_header = [c.value for c in src[1]]
old_col = {name: i + 1 for i, name in enumerate(old_header) if name}
programmes = []
for r in range(2, src.max_row + 1):
    pid = src.cell(r, old_col["ID"]).value
    if pid:
        programmes.append({"row": r, "id": str(pid), "values": {n: src.cell(r, c).value for n, c in old_col.items()}})
print(f"original programme rows: {len(programmes)}")
assert len({p["id"] for p in programmes}) == len(programmes), "duplicate IDs in source"

# ---------------------------------------------------------------- carry user edits made since, by ID
current = load_workbook(CURRENT)[SHEET]
cur_header = [c.value for c in current[1]]
cur_col = {name: i + 1 for i, name in enumerate(cur_header) if name}
carried = []
if "ID" in cur_col:
    for r in range(2, current.max_row + 1):
        ids = text(current.cell(r, cur_col["ID"]).value).split("\n")
        for column in ("Application Status", "Notes"):
            if column not in cur_col:
                continue
            cell = text(current.cell(r, cur_col[column]).value)
            if not cell:
                continue
            if len(ids) == 1 and "\n" not in cell and not re.match(r"INT-\d{6}:", cell):
                found = {ids[0]: cell}
            else:
                found = dict(re.findall(r"^(INT-\d{6}):\s*(.*)$", cell, re.MULTILINE))
            for pid, value in found.items():
                prog = next((p for p in programmes if p["id"] == pid), None)
                if prog is not None and text(prog["values"].get(column)) != value.strip():
                    carried.append((pid, column, prog["values"].get(column), value.strip()))
                    prog["values"][column] = value.strip()
print(f"user edits carried from the current workbook: {len(carried)} {carried[:5]}")

# ---------------------------------------------------------------- filter
removed, review = [], []
for prog in programmes:
    verdict, why = classify(text(prog["values"].get("Job Title")), text(prog["values"].get("Skills")))
    prog["verdict"], prog["why"] = verdict, why
    if verdict == "remove":
        removed.append(prog)
    elif verdict == "review":
        review.append(prog)
kept = [p for p in programmes if p["verdict"] != "remove"]
removed_with_user_state = [p for p in removed if text(p["values"].get("Application Status")) or text(p["values"].get("Notes"))]
print(f"non-technical removed: {len(removed)} | kept: {len(kept)} (of which for review: {len(review)})")
print(f"removed rows that carry YOUR status/notes (kept instead, listed for review): {len(removed_with_user_state)}")
for p in removed_with_user_state:
    p["verdict"], p["why"] = "review", f"would be removed ({p['why']}) but has your status/notes"
    removed.remove(p)
    review.append(p)
kept = [p for p in programmes if p["verdict"] != "remove"]

# ---------------------------------------------------------------- group consecutively by company
same_key = {company_key(k): company_key(v) for k, v in SAME_COMPANY.items()}
groups: "OrderedDict[str, list[dict]]" = OrderedDict()
for prog in kept:
    key = company_key(text(prog["values"]["Company"]))
    groups.setdefault(same_key.get(key, key), []).append(prog)


def display_name(progs):
    names = {text(p["values"]["Company"]) for p in progs}
    chosen = next((v for k, v in SAME_COMPANY.items() if k in names or v in names), None)
    if not chosen:
        trusted = [text(p["values"]["Company"]) for p in progs if text(p["values"].get("Primary Source")) not in AGGREGATORS]
        chosen = Counter(trusted or [text(p["values"]["Company"]) for p in progs]).most_common(1)[0][0]
    return re.sub(r"\b(Llp|Plc|Llc|Uk)\b", lambda m: m.group(1).upper(), chosen)


# ---------------------------------------------------------------- build the new sheet in the source workbook
new_header = COMPANY_LEVEL + [h for h in old_header if h and h not in COMPANY_LEVEL]
new_col = {name: i + 1 for i, name in enumerate(new_header)}
ws = src_wb.create_sheet("Internships (new)", 0)

for name, c in new_col.items():
    cell = ws.cell(1, c, name)
    cell._style = copy(src.cell(1, old_col[name])._style)
    ws.column_dimensions[get_column_letter(c)].width = src.column_dimensions[get_column_letter(old_col[name])].width

placement = []                      # (group key, display, first row, last row, progs)
row = 2
for key, progs in groups.items():
    name = display_name(progs)
    first = row
    for prog in progs:
        for column, c in new_col.items():
            source_cell = src.cell(prog["row"], old_col[column])
            cell = ws.cell(row, c)
            cell._style = copy(source_cell._style)
            value = name if column == "Company" else prog["values"].get(column)
            cell.value = value
            if source_cell.hyperlink:
                cell.hyperlink = copy(source_cell.hyperlink)
                cell.hyperlink.ref = cell.coordinate
            wrap_alignment = copy(cell.alignment)
            cell.alignment = Alignment(horizontal=wrap_alignment.horizontal, vertical="center", wrap_text=True)
        prog["new_row"] = row
        row += 1
    placement.append((key, name, first, row - 1, progs))
last_row = row - 1

merged_groups = 0
merges = {"Company": 0, "Match Score": 0, "Location": 0}
for key, name, first, last, progs in placement:
    if last == first:
        continue
    merged_groups += 1
    for column in COMPANY_LEVEL:
        values = [text(p["values"].get(column)) for p in progs]
        if column == "Company" or len(set(values)) == 1:
            c = new_col[column]
            ws.merge_cells(start_row=first, start_column=c, end_row=last, end_column=c)
            top = ws.cell(first, c)
            top.alignment = Alignment(horizontal=top.alignment.horizontal, vertical="center", wrap_text=True)
            merges[column] += 1

# Row heights sized to the wrapped text.
for r in range(2, last_row + 1):
    lines = 1
    for column, c in new_col.items():
        if column in COMPANY_LEVEL:
            continue
        value = text(ws.cell(r, c).value)
        if value:
            width = ws.column_dimensions[get_column_letter(c)].width or 12
            per_line = max(int(width * 1.1), 6)
            lines = max(lines, sum(max(1, -(-len(part) // per_line)) for part in value.split("\n")))
    ws.row_dimensions[r].height = min(MAX_ROW_HEIGHT, LINE_HEIGHT * lines)

last_letter = get_column_letter(len(new_header))
ws.auto_filter.ref = f"A1:{last_letter}{last_row}"
ws.freeze_panes = src.freeze_panes and "D2"


def remap(sqref: str) -> str:
    out = []
    for rng in str(sqref).split():
        m = re.match(r"^\$?([A-Z]+)\$?\d+(?::\$?([A-Z]+)\$?\d+)?$", rng)
        if not m:
            continue
        from openpyxl.utils import column_index_from_string
        left = new_col[old_header[column_index_from_string(m.group(1)) - 1]]
        right = new_col[old_header[column_index_from_string(m.group(2) or m.group(1)) - 1]]
        out.append(f"{get_column_letter(left)}2:{get_column_letter(right)}{last_row}")
    return " ".join(out)


for dv in src.data_validations.dataValidation:
    new_dv = DataValidation(type=dv.type, formula1=dv.formula1, formula2=dv.formula2, allow_blank=dv.allow_blank,
                            showDropDown=dv.showDropDown, showErrorMessage=dv.showErrorMessage,
                            showInputMessage=dv.showInputMessage, operator=dv.operator)
    new_dv.sqref = remap(dv.sqref)
    if str(new_dv.sqref):
        ws.add_data_validation(new_dv)
for cf in src.conditional_formatting:
    target = remap(str(cf.sqref))
    for rule in cf.rules:
        if target:
            ws.conditional_formatting.add(target, rule)

del src_wb[SHEET]
ws.title = SHEET
src_wb.active = src_wb[SHEET]

# ---------------------------------------------------------------- other sheets drop the same programmes
removed_ids = {p["id"] for p in removed}
other_sheet_rows = {}
for sheet_name in src_wb.sheetnames:
    if sheet_name == SHEET:
        continue
    other = src_wb[sheet_name]
    header = [c.value for c in other[1]]
    if "ID" in header:
        id_col = header.index("ID") + 1
        doomed = [r for r in range(2, other.max_row + 1) if text(other.cell(r, id_col).value) in removed_ids]
    elif "Company" in header and "Job Title" in header:
        kept_pairs = {(text(p["values"]["Company"]), text(p["values"]["Job Title"])) for p in kept}
        removed_pairs = {(text(p["values"]["Company"]), text(p["values"]["Job Title"])) for p in removed} - kept_pairs
        cc, tc = header.index("Company") + 1, header.index("Job Title") + 1
        doomed = [r for r in range(2, other.max_row + 1)
                  if (text(other.cell(r, cc).value), text(other.cell(r, tc).value)) in removed_pairs]
    else:
        continue
    for r in reversed(doomed):
        other.delete_rows(r)
    other._hyperlinks = [h for h in other._hyperlinks if h.ref and int(re.sub(r"\D", "", h.ref) or 0) <= other.max_row]
    if other.auto_filter.ref:
        other.auto_filter.ref = f"A1:{get_column_letter(other.max_column)}{other.max_row}"
    other_sheet_rows[sheet_name] = len(doomed)

src_wb.save(DRAFT)
print(f"draft written: {DRAFT}")

# ---------------------------------------------------------------- verify from disk
v = load_workbook(DRAFT)[SHEET]
problems = []
header_now = [c.value for c in v[1]]
if header_now != new_header:
    problems.append(("header", header_now))
ids_now = Counter(text(v.cell(r, new_col["ID"]).value) for r in range(2, v.max_row + 1) if v.cell(r, new_col["ID"]).value)
kept_ids = {p["id"] for p in kept}
if set(ids_now) != kept_ids:
    problems.append(("ids", len(set(ids_now) ^ kept_ids)))
if any(c > 1 for c in ids_now.values()):
    problems.append(("id twice", [i for i, c in ids_now.items() if c > 1][:5]))
merged_ranges = [str(m) for m in v.merged_cells.ranges]
programme_letters = {get_column_letter(c) for name, c in new_col.items() if name not in COMPANY_LEVEL}
for m in merged_ranges:
    if re.match(r"^([A-Z]+)", m).group(1) in programme_letters:
        problems.append(("programme column merged", m))
merged_lookup = {}
for m in v.merged_cells.ranges:
    for r in range(m.min_row, m.max_row + 1):
        merged_lookup[(r, m.min_col)] = (m.min_row, m.max_row)
for key, name, first, last, progs in placement:
    for column in COMPANY_LEVEL:
        c = new_col[column]
        span = merged_lookup.get((first, c))
        values = [text(p["values"].get(column)) for p in progs]
        should = last > first and (column == "Company" or len(set(values)) == 1)
        if should and span != (first, last):
            problems.append(("missing merge", name, column))
        if not should and span:
            problems.append(("merged although different", name, column, values))
    for prog in progs:
        r = prog["new_row"]
        for column, c in new_col.items():
            cell = v.cell(r, c)
            if column in COMPANY_LEVEL:
                span = merged_lookup.get((r, c))
                shown = text(v.cell(span[0], c).value) if span else text(cell.value)
                expected = name if column == "Company" else text(prog["values"].get(column))
                if shown != expected:
                    problems.append((prog["id"], column, expected, shown))
                continue
            if text(cell.value) != text(prog["values"].get(column)):
                problems.append((prog["id"], column, prog["values"].get(column), cell.value))
leftover = [(text(v.cell(r, new_col["ID"]).value), text(v.cell(r, new_col["Job Title"]).value))
            for r in range(2, v.max_row + 1)
            if classify(text(v.cell(r, new_col["Job Title"]).value), text(v.cell(r, new_col["Skills"]).value))[0] == "remove"
            and text(v.cell(r, new_col["ID"]).value) not in {p["id"] for p in review}]
if leftover:
    problems.append(("non-technical still present", leftover[:5]))
if v.auto_filter.ref != f"A1:{last_letter}{last_row}":
    problems.append(("filter", v.auto_filter.ref))

# Links: every retained programme that had a clickable link still has the same one.
source_again = load_workbook(SOURCE)[SHEET]
for prog in kept:
    old_cell = source_again.cell(prog["row"], old_col["Application URL"])
    new_cell = v.cell(prog["new_row"], new_col["Application URL"])
    if old_cell.hyperlink and (not new_cell.hyperlink or new_cell.hyperlink.target != old_cell.hyperlink.target):
        problems.append((prog["id"], "Application URL link", old_cell.hyperlink.target, new_cell.hyperlink))

# ---------------------------------------------------------------- report + lists
with REMOVED_CSV.open("w", newline="", encoding="utf-8-sig") as fh:
    w = csv.writer(fh)
    w.writerow(["ID", "Company", "Job Title", "Why removed", "Eligible for You", "Deadline"])
    for p in removed:
        w.writerow([p["id"], text(p["values"]["Company"]), text(p["values"]["Job Title"]), p["why"],
                    text(p["values"].get("Eligible for You")), text(p["values"].get("Deadline"))[:10]])
with REVIEW_CSV.open("w", newline="", encoding="utf-8-sig") as fh:
    w = csv.writer(fh)
    w.writerow(["ID", "Company", "Job Title", "Why it needs a look", "Eligible for You", "Deadline"])
    for p in review:
        w.writerow([p["id"], text(p["values"]["Company"]), text(p["values"]["Job Title"]), p["why"],
                    text(p["values"].get("Eligible for You")), text(p["values"].get("Deadline"))[:10]])

print(f"technical programme rows remaining: {len(kept)}")
print(f"companies represented: {len(placement)} | company groups with merged cells: {merged_groups}")
print(f"merged: Company {merges['Company']}, Match Score {merges['Match Score']}, Location {merges['Location']}")
print(f"other sheets, rows removed for the same programmes: {other_sheet_rows}")
print(f"removed that were Eligible=Yes for you: "
      f"{[(p['id'], text(p['values']['Company']), text(p['values']['Job Title'])[:50]) for p in removed if text(p['values'].get('Eligible for You')) == 'Yes']}")
print(f"review reasons: {Counter(p['why'].split(':')[0].split(' (')[0] for p in review).most_common()}")
print(f"verification problems: {len(problems)}")
for p in problems[:25]:
    print("   ", p)

if problems:
    print("VERIFICATION FAILED - the workbook was NOT replaced.")
    sys.exit(2)
if "--check" in sys.argv:
    print("verification passed (--check: workbook not replaced)")
    sys.exit(0)
if (ROOT / "data" / "~$internships.xlsx").exists():
    print("Excel has the workbook open - close it; the verified draft is at", DRAFT)
    sys.exit(3)
shutil.copy2(CURRENT, BACKUP)
shutil.copy2(DRAFT, CURRENT)
print(f"verified. backup of the previous workbook: {BACKUP} | replaced: {CURRENT}")

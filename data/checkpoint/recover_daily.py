"""Put back the rows the 17 Sep daily scan lost to failed sources.

Gradcracker and Bright Network failed (browser closed) and had no snapshot, since
their rows came from pages the user saved; Trackr's three 2026-season slices timed
out. A failed source should degrade, not lose rows, so:

* the Trackr 2026 slices are fetched again (3 API calls, nothing else is searched);
* Gradcracker / Bright Network rows from the last complete state are carried forward
  and saved as those sources' snapshots, so a later failure reuses them.
Backs up state and workbook first. Writes through excel.write, which keeps statuses.
"""
import json
import os
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

from openpyxl import load_workbook  # noqa: E402

from sources import trackr  # noqa: E402
from sources.fetch import Session, SourceError  # noqa: E402
from tracker import config, dedupe, eligibility, excel, graduation, pipeline, programme  # noqa: E402
from tracker.models import Job  # noqa: E402

BACKUPS = ROOT / "data" / "checkpoint" / "backups"
if (ROOT / "data" / "~$internships.xlsx").exists():
    sys.exit("Excel has the workbook open - close it first.")
for src, name in [(config.STATE_DIR / "last_run.json", "last_run-daily-0917.json"),
                  (config.STATE_DIR / "id_registry.json", "id_registry-daily-0917.json"),
                  (config.WORKBOOK, "internships-daily-0917-formatted.xlsx")]:
    shutil.copy2(src, BACKUPS / name)
for snap in (config.STATE_DIR / "snapshots").glob("*.json"):
    shutil.copy2(snap, BACKUPS / f"snapshot-0917-{snap.name}")

current = [Job.from_dict(d) for d in json.loads((config.STATE_DIR / "last_run.json").read_text(encoding="utf-8"))["jobs"]]
old = [Job.from_dict(d) for d in json.loads((BACKUPS / "last_run-before-remerge.json").read_text(encoding="utf-8"))["jobs"]]
ws = load_workbook(BACKUPS / "internships-before-company-consolidation.xlsx", read_only=True)["Internships"]
rows = ws.iter_rows(values_only=True)
id_col = next(rows).index("ID")
complete_ids = {r[id_col] for r in rows if r[id_col]}
old = [j for j in old if j.id in complete_ids]
current_ids = {j.id for j in current}
profile = config.load_profile()

# ---- Trackr 2026 slices
settings = dict(config.load("sources")["sources"]["trackr"], seasons=["2026"])
try:
    with Session() as session:
        refetched = trackr.search_jobs(session, settings)
    print(f"Trackr 2026 slices: {len(refetched)} listings")
except SourceError as exc:
    refetched = []
    print("Trackr 2026 slices failed again:", exc)
for job in refetched:
    programme.refine(job, f"{job.year_of_study}. {job.description}")
    graduation.apply(job, f"{job.year_of_study}. {job.description}", profile)
    eligibility.assess(job, job.description, profile)

# ---- saved-page sources
FAILED = {"gradcracker", "brightnetwork"}
carried = [j for j in old if j.id not in current_ids and {r.source for r in j.sources} & FAILED]
print(f"Gradcracker/Bright Network rows carried forward: {len(carried)}")

merged, dupes = dedupe.deduplicate(current + refetched + carried)
dedupe.assign_ids(merged)
added = [j for j in merged if j.id not in current_ids]
pipeline._restore_state(added, pipeline._match_previous(added, old))
pipeline._score_all(merged, profile, config.load("search_filters"), config.load("matching_weights"))
print(f"{len(current)} + {len(refetched)} refetched + {len(carried)} carried - {dupes} duplicates = {len(merged)} rows "
      f"({len(added)} restored)")

written = excel.write(merged)
if written != config.WORKBOOK:
    sys.exit(f"workbook was open; wrote {written} - state NOT saved")
pipeline._save_previous(merged)
if refetched:
    trackr_now = [j for j in current if any(r.source == "trackr" for r in j.sources)]
    pipeline._save_snapshot("trackr", pipeline._load_snapshot("trackr") + refetched)
for source in FAILED:
    rows_for = [j for j in merged if any(r.source == source for r in j.sources)]
    pipeline._save_snapshot(source, rows_for)
    print(f"snapshot {source}: {len(rows_for)}")
print("workbook written:", written)
